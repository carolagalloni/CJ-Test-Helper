"""Unit tests for cjhelper.outputs (spec, "Outputs" and the file format of "Data generation rules").

The inputs come from the hand-made world of tests/test_datagen.py, run through datagen.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import replace
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook

from cjhelper.datagen import SCENARIO_COLUMNS
from cjhelper.model import (
    ENTRY_OTHER,
    ENTRY_RUN_ONCE,
    ENTRY_SCHEDULED,
    GENERATED_KEY,
    KIND_SMS,
    KIND_WAIT_ATTRIBUTE,
    KIND_WAIT_DURATION,
    KIND_UPDATE,
    ROLE_LINK,
    STATUS_COVERED,
    ActionWindow,
    Activity,
    BranchStatus,
    CheckResult,
    Field,
    GenerationResult,
    JourneyFilter,
    Outcome,
    Relationship,
    Table,
    TableData,
    TimelinePoint,
    Write,
)
from cjhelper import outputs
from cjhelper.outputs import (
    REPORT_FILE,
    SCENARIOS_FILE,
    TESTBOOK_COLUMNS,
    TESTBOOK_FILE,
    build_zip,
    read_csv,
    table_file_name,
    build_testbook_rows,
    unsupported_warning,
    write_report,
    write_scenarios_csv,
    write_table_csv,
    write_testbook,
)
from tests.test_datagen import CONFIG, ENTRY, MASTER, build_world, run


def build_result(checks=None, **settings) -> GenerationResult:
    """A full GenerationResult, with every file of the zip written as the pipeline would."""
    world = build_world(**settings)
    data = run(world)
    result = GenerationResult(
        package=world.package, table_model=world.model, coverage=world.coverage, data=data,
        checks=checks if checks is not None else [
            CheckResult("Values respect the schema", True, []),
            CheckResult("Every key joins", False, ["DE_Master row 4 has no scenario row.", "Second detail."]),
        ],
        settings=world.settings, files={}, generation_time=world.generation_time)
    for name, table_data in data.tables.items():
        result.files[table_file_name(name)] = write_table_csv(table_data, world.package.tables[name])
    result.files[SCENARIOS_FILE] = write_scenarios_csv(data.scenarios_rows)
    result.files[REPORT_FILE] = write_report(result).encode("utf-8")
    result.files[TESTBOOK_FILE] = write_testbook(result)
    return result


def small_table() -> Table:
    # Fields listed out of schema order on purpose: the header follows the ordinals.
    return Table(id="t", name="DE_Small", customer_key=None, fields=[
        Field("Flag", "Boolean", None, False, True, None, 2),
        Field("Id", "Text", 10, True, False, None, 0),
        Field("Name", "Text", 50, False, True, None, 1),
        Field("Note", "Text", 50, False, True, None, 3),
    ])


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------

def test_table_csv_format():
    data = TableData("DE_Small", ["Id", "Name", "Flag", "Note"], [
        {"Id": "1", "Name": "Caffè, \"bar\"", "Flag": "True", "Note": None},
        {"id": "2", "NAME": "plain", "flag": "False", "note": "line\nbreak"},   # other letter case: still found
    ])
    content = write_table_csv(data, small_table())
    assert not content.startswith(b"\xef\xbb\xbf")                 # no byte order mark
    assert b"\r\n" not in content.replace(b"line\nbreak", b"")     # "\n" line ends
    assert content.decode("utf-8") == (
        "Id,Name,Flag,Note\n"
        "1,\"Caffè, \"\"bar\"\"\",True,\n"
        "2,plain,False,\"line\nbreak\"\n"
    )


def test_table_csv_header_is_the_schema_even_for_missing_values():
    content = write_table_csv(TableData("DE_Small", [], [{"Id": "1"}]), small_table())
    assert content == b"Id,Name,Flag,Note\n1,,,\n"


def test_read_csv_round_trip():
    header, rows = read_csv(b"A,B,C\n1,,\"x, y\"\n\"multi\nline\",2,\n")
    assert header == ["A", "B", "C"]
    assert rows == [{"A": "1", "B": None, "C": "x, y"}, {"A": "multi\nline", "B": "2", "C": None}]
    assert read_csv(b"") == ([], [])
    assert read_csv(b"A\n\n1\n") == (["A"], [{"A": "1"}])                 # blank line skipped
    assert read_csv(write_table_csv(TableData("T", [], [{"Id": None}]), small_table()))[1] == [
        {"Id": None, "Name": None, "Flag": None, "Note": None}]


def test_generated_csvs_read_back_identical():
    result = build_result()
    for name, table_data in result.data.tables.items():
        header, rows = read_csv(result.files[table_file_name(name)])
        assert header == result.package.tables[name].field_names
        assert rows == [{column: (value if value != "" else None) for column, value in row.items()}
                        for row in table_data.rows]


def test_generated_csv_writes_booleans_and_nulls_as_the_spec_says():
    result = build_result()
    text = result.files[table_file_name(CONFIG)].decode("utf-8")
    lines = text.split("\n")
    assert lines[0] == "CampaignName,CHANNEL,ACTIVE,HAS_RUN,START_DATE,END_DATE,RUN_TIME,Region,WAIT1,Label"
    assert lines[1] == ("TST_R01_S01,Email,True,False,2026-10-10 09:30,2026-10-18 09:30,,North,"
                        "2026-10-12 09:30,TST_La")
    assert text.endswith("\n") and len(lines) == 4                 # header, 2 rows, final empty string


def test_scenarios_csv():
    result = build_result()
    header, rows = read_csv(result.files[SCENARIOS_FILE])
    assert header == SCENARIO_COLUMNS + [f"{CONFIG}.CHANNEL"]
    assert [r["scenario_key"] for r in rows] == ["TST_R01_S01", "TST_R01_S02"]
    assert [r["active_at_start"] for r in rows] == ["yes", "no"]
    assert rows[1]["contact_count"] == "50"


def test_scenarios_csv_without_scenarios_still_has_a_header():
    assert write_scenarios_csv([]) == (",".join(SCENARIO_COLUMNS) + "\n").encode("utf-8")


# ---------------------------------------------------------------------------
# Zip
# ---------------------------------------------------------------------------

def test_zip_contains_every_file():
    result = build_result()
    archive = zipfile.ZipFile(io.BytesIO(build_zip(result.files)))
    assert archive.namelist() == list(result.files)
    assert archive.namelist() == [f"{ENTRY}.csv", f"{MASTER}.csv", f"{CONFIG}.csv", SCENARIOS_FILE, REPORT_FILE,
                                  TESTBOOK_FILE]
    for name, content in result.files.items():
        assert archive.read(name) == content


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------

def test_report_sections_in_order():
    report = write_report(build_result())
    titles = [line for line in report.splitlines() if line.startswith("## ")]
    assert titles == ["## Riepilogo", "## Trust check", "## Tabelle", "## Scenari", "## Branch non coperti",
                      "## Branch coperti con un'azione manuale o tramite un random split",
                      "## Elementi non supportati", "## Warning e pattern sospetti", "## Assunzioni",
                      "## Impostazioni usate"]


def test_report_totals_match_the_data():
    result = build_result()
    report = write_report(result)
    assert "I dati di test comprendono 2 scenari, 3 casi e 56 contatti." in report
    assert "Su 10 branch, 8 sono coperti, 1 è coperto con un'azione manuale e 1 non è coperto." in report
    assert "| Contatti | 56 |" in report
    assert "| Contatti con l'inbox reale | 3 |" in report
    assert "| Branch non coperti | 1 |" in report
    # rows per CSV match the files
    for name in result.data.tables:
        _, rows = read_csv(result.files[table_file_name(name)])
        assert f"| {table_file_name(name)} | {len(rows)} |" in report
    assert "| DE_Log | written | - | no | - |" in report


def test_report_lists_checks_branches_unsupported_warnings_and_settings():
    result = build_result()
    promo = next(b for b in result.coverage.branches if b.status == "not covered")
    promo.reason = "non raggiungibile: l'email di riferimento 'Promo' (EMAIL-2) non contiene link da cliccare"   # graph.py
    report = write_report(result)
    assert "1 trust check su 2 è superato, 1 non è superato: i dettagli sono qui sotto." in report
    assert "| Every key joins | NON SUPERATO | 2 |" in report
    assert "- DE_Master row 4 has no scenario row." in report
    assert ("| Clicked Promo (ENG-2) | Yes (yes) | Non raggiungibile perché l'email di riferimento 'Promo' (EMAIL-2) "
            "non contiene link da cliccare. |") in report
    assert "`C.Region LIKE 'N%' OR C.Region IS NULL`" in report
    assert "Placeholder: DE_Entry.Weird" in report and "Placeholder: DE_Entry.Mandatory" in report
    for warning in ("Reading warning from the loader.", "Table selection warning.",
                    "Parser warning about the entry query.", "Coverage warning."):
        assert f"- {warning}" in report
    assert "| Inbox reale per i test di open e click | tester@example.com |" in report
    assert "| Un solo scenario attivo alla volta, come applicato | sì, campo di attivazione DE_Config.ACTIVE |" in report
    assert "| Ora di generazione | 2026-10-11 09:30 |" in report
    assert "W. Europe Standard Time" in report


def test_report_describes_the_direct_import_only():
    report = write_report(build_result())
    how = report[report.index("### Come eseguire i test"):report.index("### Elenco degli scenari")]
    assert how.startswith("### Come eseguire i test\n\nImporta prima DE_Master.csv e DE_Config.csv, poi DE_Entry.csv,")
    assert how.count("automation") == 3 and how.count("A_Entry") == 1    # one note on the entry source only
    assert "Assunzione da verificare al primo run reale: l'ingresso in questo journey è pianificato " \
           "dall'automation A_Entry." in how
    assert "ACTIVE è un filtro della entry query e nessuno split lo legge, quindi con l'import diretto non " \
           "trattiene alcuno scenario." in how
    for alternative in ("procedura", "Tramite l'automation", "Esegui l'automation", "Run Once"):
        assert alternative not in report
    assumptions = report[report.index("## Assunzioni"):report.index("## Impostazioni usate")]
    assert "- L'import diretto di DE_Entry.csv fa entrare i suoi contatti." in assumptions


def test_report_scenarios_and_cases():
    report = write_report(build_result())
    assert ("| TST_R01_S01 | 1 | sì | True | False | 2 | 6 | da TST_R01_C0001 a TST_R01_C0006 | "
            "DE_Config.CHANNEL = Email |") in report
    assert "| S02-C01 | TST_R01_S02 | 50 | da TST_R01_C0007 a TST_R01_C0056 |" in report
    assert "da 1 a 9 contatti attesi sul branch 10%" in report


def test_report_without_checks_and_without_problems():
    result = build_result(checks=[])
    result.coverage.branches = [b for b in result.coverage.branches if b.status != "not covered"]
    report = write_report(result)
    assert "Su questi file non è stato eseguito alcun trust check." in report
    assert "Tutti i branch sono coperti." in report


def test_suspicious_patterns_are_gathered_from_every_module():
    result = build_result()
    # The texts as coverage.py and loader.py write them.
    result.coverage.warnings += [
        "pattern sospetto: DE_Config.CHANNEL è confrontato con 'Email' e 'EMAIL'; sono valori diversi perché i "
        "confronti di testo distinguono maiuscole e minuscole (impostazione)",
        "pattern sospetto: 2 split con condizioni diverse hanno lo stesso nome 'Channel' (DEC-9, DEC-1)",
    ]
    result.package.warnings += [
        "pattern sospetto: la costante 'X' è scritta nel campo di chiave primaria DE_Log.Id da 2 attività Update "
        "contact, quindi tutti i contatti hanno lo stesso valore di chiave in quel campo (UPD-2, UPD-1).",
        "pattern sospetto: la relazione uno a uno 'M -> E' collega DE_Entry tramite ContactId, ma la chiave primaria "
        "della tabella comprende anche ActionId, CampaignName: a un contatto possono corrispondere più righe.",
        "Pattern sospetto: due split hanno lo stesso nome 'Channel'.",
    ]
    report = write_report(result)
    section = report[report.index("### Pattern sospetti"):report.index("### Lettura dell'export")]
    assert section.count("\n- ") == 5
    assert ("- DE_Config.CHANNEL viene confrontato con 'Email' e 'EMAIL'. Con l'impostazione attuale i confronti di "
            "testo distinguono maiuscole e minuscole, quindi sono valori diversi. Come li confronta Salesforce è "
            "un'assunzione da verificare.") in section
    assert ("- Alcuni split con lo stesso nome verificano condizioni diverse: 'Channel' in 2 versioni, comparse per la "
            "prima volta in DEC-1 e DEC-9.") in section
    assert ("- 2 attività Update contact scrivono la costante 'X' nel campo di chiave primaria DE_Log.Id. Quindi tutti "
            "i contatti hanno lo stesso valore in quel campo. Le attività sono UPD-1 e UPD-2.") in section
    assert ("- Una relazione uno a uno collega una tabella su una parte della sua chiave primaria, quindi a un "
            "contatto possono corrispondere più righe: 'M -> E' collega DE_Entry tramite ContactId, ma la chiave "
            "primaria di DE_Entry comprende anche ActionId e CampaignName.") in section
    assert "- Due split hanno lo stesso nome 'Channel'." in section
    assert "pattern sospetto:" not in report.lower()             # the prefix is dropped once the line is placed
    coverage = report[report.index("### Coverage"):report.index("### Generazione dei dati")]
    assert "confrontato" not in coverage and "stesso nome" not in coverage


def test_not_covered_reasons_from_coverage_are_reworded_and_grouped():
    result = build_result()
    by_ref = {(b.activity_key, b.outcome_key): b for b in result.coverage.branches}
    root_reason = ("non raggiungibile: lo split è raggiunto, ma su quei percorsi la condizione di 'Has permission' "
                   "non è mai vera (DE_Master.Perm vale sempre uno tra 'False', null)")
    root = by_ref[("DEC-2", "perm")]
    root.status, root.reason, root.covered_by = "not covered", root_reason, []
    for outcome in ("yes", "no"):
        branch = by_ref[("ENG-2", outcome)]
        branch.status, branch.covered_by, branch.behind = "not covered", [], ("DEC-2", "perm")
        branch.reason = ("non raggiungibile: si trova a valle del branch 'Permission: Has permission' (DEC-2), che "
                         "non è coperto per questo motivo: " + root_reason)
    after = by_ref[("DEC-1", "remainder_path")]
    after.status, after.covered_by = "not covered", []
    after.reason = ("a valle di un elemento non supportato: questo branch viene dopo il branch non supportato 'Email' "
                    "dello stesso split, quindi non viene mai valutato")
    rb5 = by_ref[("ENG-1", "no")]
    rb5.status, rb5.covered_by = "not covered", []
    rb5.reason = ("branch non supportato: lo split misura un tipo di statistica che non è né un open né un click, "
                  "quindi nessuno dei suoi branch può essere considerato coperto")
    report = write_report(result)
    section = report[report.index("## Branch non coperti"):report.index("## Branch coperti con")]
    assert ("5 branch non sono coperti. Di questi, 3 si trovano a valle di 1 branch che nessun contatto può "
            "raggiungere o 1 elemento non supportato: le cause sono spiegate per prime. I restanti 2 hanno ciascuno "
            "un motivo proprio, indicato nella tabella alla fine.") in section
    assert "### Cause a monte" in section
    assert ("2 branch si trovano a valle di 'Permission: Has permission' (DEC-2), che non è raggiungibile perché la "
            "sua condizione") in section
    assert "dove lo split viene raggiunto, Perm vale sempre 'False' o vuoto." in section
    assert ("1 branch viene dopo il branch non supportato 'Email' dello stesso split, quindi non viene mai valutato. "
            "Vedi la sezione Elementi non supportati.") in section
    assert ("| Opened Welcome (ENG-1) | No (no) | Non coperto perché il branch non è supportato: lo split misura un "
            "tipo di statistica che non è né un open né un click, quindi nessuno dei suoi branch può essere "
            "considerato coperto. |") in section
    assert "A valle di questo branch si trovano 2 branch." in section


# ---------------------------------------------------------------------------
# Testbook
# ---------------------------------------------------------------------------

def test_testbook_has_one_row_per_case_and_names_every_contact():
    result = build_result()
    rows = build_testbook_rows(result)
    assert len(rows) == len(result.coverage.cases) == 3
    assert [row[0] for row in rows] == [1, 2, 3]
    assert {row[1] for row in rows} == {"SFMC"} and {row[7] for row in rows} == {"TO START"}
    assert [row[2] for row in rows] == ["S01-C01 - Welcome opened", "S01-C02 - Percorso fino a Wait one minute",
                                        "S02-C01 - Random split 90/10"]
    text = "\n".join(str(value) for row in rows for value in row)
    for contact in result.data.contacts:
        if contact.case_id != "S02-C01":
            assert contact.key in text
    assert "Contatti di test (50): da TST_R01_C0007 a TST_R01_C0056." in rows[2][4]
    random_keys = [c.key for c in result.data.contacts if c.case_id == "S02-C01"]
    assert random_keys == [f"TST_R01_C{n:04d}" for n in range(7, 57)]   # the range names exactly these


def test_testbook_preconditions_steps_and_expected_result():
    rows = build_testbook_rows(build_result())
    first = dict(zip(TESTBOOK_COLUMNS, rows[0]))
    assert ("Importa questi file CSV: DE_Entry.csv, DE_Master.csv, DE_Config.csv. Importa DE_Entry.csv per ultimo."
            in first["Pre conditions"])
    assert "Scenario attivo: TST_R01_S01" in first["Pre conditions"]
    assert "Percorso atteso: Channel [Email] > Welcome > Wait campaign date > Wait reminder date > " \
           "Opened Welcome [Yes] > Log open." in first["Description"]
    steps = first["Detailed steps"].split("\n")
    assert [step.split(".")[0] for step in steps] == ["1", "2", "3", "4", "5"]
    assert steps[1] == ("2. Avvia l'ingresso nel journey Test journey dalle righe importate di DE_Entry.csv. Tutti i "
                        "test di questo scenario partono insieme.")
    text = "\n".join(str(value) for row in rows for value in row)
    assert "automation" not in text and "oppure" not in text            # direct import only, no alternative
    assert "Prima del 2026-10-12 09:30, apri l'email 'Welcome'" in steps[3]
    assert "tester+TST_R01_C0001@example.com" in steps[3]
    assert "- Email ricevute: Welcome." in first["Expected result"]
    assert "- Valori scritti dal journey: DE_Log.Status = Opened." in first["Expected result"]
    assert "- Ultima attività raggiunta: Log open (UPD-1)." in first["Expected result"]
    assert first["Note"].startswith("Azione manuale richiesta: aprire l'email 'Welcome' prima del 2026-10-12 09:30.")

    second = dict(zip(TESTBOOK_COLUMNS, rows[1]))
    assert "Profilo del contatto: DE_Master.Perm = True." in second["Description"]     # CSV form of "true"
    assert "Wait one minute (1 minuto)" in second["Detailed steps"]
    assert second["Note"] == ""

    third = dict(zip(TESTBOOK_COLUMNS, rows[2]))
    assert "attiva lo scenario TST_R01_S02: in DE_Config, imposta ACTIVE a True sulla riga con CampaignName " \
           "TST_R01_S02 e a False su tutte le altre righe." in third["Detailed steps"]
    assert "- 10%: email ricevute: Promo; valori scritti: nessuno; ultima attività raggiunta: Clicked Promo (ENG-2)." \
           in third["Expected result"]
    assert "Intervallo accettato: da 1 a 9 contatti sul branch 10%." in third["Expected result"]
    assert "Sul branch 10% sono attesi da 1 a 9 dei 50 contatti." in third["Note"]


def test_testbook_without_one_active_scenario_or_real_inbox():
    rows = build_testbook_rows(build_result(one_active_scenario=False, real_inbox=""))
    first = dict(zip(TESTBOOK_COLUMNS, rows[0]))
    assert "Tutti gli scenari sono attivi contemporaneamente." in first["Pre conditions"]
    assert "non è stata indicata un'inbox reale" in first["Detailed steps"]


def test_testbook_xlsx_layout():
    result = build_result()
    sheet = load_workbook(io.BytesIO(result.files[TESTBOOK_FILE])).active
    values = list(sheet.iter_rows(values_only=True))
    assert list(values[0]) == TESTBOOK_COLUMNS
    assert len(values) == 1 + len(result.coverage.cases)
    assert [list(row) for row in values[1:]] == [[v if v != "" else None for v in row] for row in build_testbook_rows(result)]


# ---------------------------------------------------------------------------
# Determinism
# ---------------------------------------------------------------------------

def test_two_runs_give_identical_bytes():
    first, second = build_result(), build_result()
    assert first.files == second.files
    assert build_zip(first.files) == build_zip(second.files)
    assert write_testbook(first) == write_testbook(first)


# ---------------------------------------------------------------------------
# Spec, "Extensions beyond the pilot": entry procedures, rows absent, honest outcomes
# ---------------------------------------------------------------------------

def _how_to_run(report: str) -> str:
    return report[report.index("### Come eseguire i test"):report.index("### Elenco degli scenari")]


def _rows(result) -> list[dict]:
    return [dict(zip(TESTBOOK_COLUMNS, row)) for row in build_testbook_rows(result)]


def _with_entry(kind, event_type="EmailAudience", **settings):
    """The hand-made world with another entry source: no feeding automation, as loader.py reads it."""
    result = build_result(**settings)
    result.package.entry_event.kind = kind
    result.package.entry_event.event_type = event_type
    result.package.entry_event.automation_name = None
    return result


def test_run_once_procedure_in_report_and_testbook():
    result = _with_entry(ENTRY_RUN_ONCE, one_active_scenario=False)
    report = write_report(result)
    how = _how_to_run(report)
    assert ("L'entry source di questo journey è la data extension DE_Entry con pianificazione Run Once. Importa prima "
            "DE_Master.csv e DE_Config.csv, poi DE_Entry.csv, e solo dopo attiva il journey. Entrano solo le righe "
            "presenti nella data extension quando il journey viene attivato: le righe importate dopo non entrano "
            "mai.") in how
    assert "Quando attivi il journey, DE_Entry deve contenere solo righe di test" in how
    assert "Un nuovo tentativo richiede una nuova versione del journey e un nuovo run identifier" in how
    assert "con Run Once non si può eseguire uno scenario alla volta" in how
    assert "automation A_Entry" not in report
    first = _rows(result)[0]
    assert "Importa tutti i file prima di attivare il journey." in first["Pre conditions"]
    assert "Tutti gli scenari entrano insieme quando il journey viene attivato." in first["Pre conditions"]
    assert ("Ogni nuovo tentativo richiede una nuova versione del journey e un nuovo run identifier."
            in first["Pre conditions"])
    steps = first["Detailed steps"].split("\n")
    assert steps[0] == ("1. Verifica che le righe dello scenario TST_R01_S01 siano state importate e che DE_Entry "
                        "contenga solo righe di test.")
    assert steps[1] == ("2. Attiva il journey Test journey: entrano solo le righe presenti in DE_Entry in quel momento, "
                        "e le righe importate dopo non entrano. Tutti i test partono insieme.")
    assumptions = report[report.index("## Assunzioni"):report.index("## Impostazioni usate")]
    assert "è lo scheduler che Journey Builder crea per la pianificazione della data extension" in assumptions


def test_run_once_ignores_one_active_scenario():
    """Spec, Run Once row: "All scenarios enter together, so 'one active scenario at a time' is not offered": the
    setting is ignored, the report gives no row-switching procedure and says why in the settings table."""
    report = write_report(_with_entry(ENTRY_RUN_ONCE))
    how = _how_to_run(report)
    assert ("Tutti gli scenari entrano insieme quando il journey viene attivato, quindi vengono eseguiti insieme: con "
            "Run Once non si può eseguire uno scenario alla volta.") in how
    assert "rendi attivo lo scenario successivo" not in report and "sì, campo di attivazione" not in report
    assert ("| Un solo scenario attivo alla volta, come applicato | no: non si applica a un ingresso Run Once, dove "
            "tutti gli scenari entrano insieme |") in report


def test_other_entry_sources_use_a_copy_of_the_journey():
    result = _with_entry(ENTRY_OTHER, "CloudPagesSmartCaptureSubmissionEvent", one_active_scenario=False)
    how = _how_to_run(write_report(result))
    assert ("L'entry source di questo journey è un form CloudPages (CloudPagesSmartCaptureSubmissionEvent): i contatti "
            "non entrano con l'import di un CSV. Il test si fa su una copia del journey la cui entry source è una data "
            "extension con le stesse colonne di DE_Entry. Importa prima DE_Master.csv e DE_Config.csv nelle loro data "
            "extension, poi DE_Entry.csv nella data extension di ingresso della copia.") in how
    first = _rows(result)[0]
    assert first["Pre conditions"].startswith(
        "Importa DE_Master.csv e DE_Config.csv, poi importa DE_Entry.csv nella data extension di ingresso della copia "
        "del journey, che ha le stesse colonne di DE_Entry.")
    assert "2. Fai entrare i contatti nella copia del journey Test journey" in first["Detailed steps"]
    api = _with_entry(ENTRY_OTHER, "APIEvent", one_active_scenario=False)
    assert "è un API event (APIEvent)" in write_report(api)


def test_recurring_schedule_is_a_procedure_to_verify():
    result = _with_entry(ENTRY_SCHEDULED, one_active_scenario=False)
    result.package.entry_event.schedule_mode = "recurring"
    report = write_report(result)
    assert ("L'entry source di questo journey è la data extension DE_Entry con una pianificazione ricorrente "
            "(recurring). Importa prima DE_Master.csv e DE_Config.csv, poi DE_Entry.csv. Procedura da verificare"
            in _how_to_run(report))
    assert "entrano alla valutazione successiva della data extension di ingresso" in report
    assert "(procedura da verificare)" in _rows(result)[0]["Detailed steps"]


@pytest.mark.parametrize("mode", ["OnceAndDone", "SingleEntryAcrossAllVersions"])
def test_single_entry_modes_need_a_new_run_identifier(mode):
    result = build_result()
    result.package.journey.entry_mode = mode
    report = write_report(result)
    assert f"La modalità di ingresso del journey è {mode}" in _how_to_run(report)
    assert (f"Con la modalità di ingresso {mode}, i contatti di un tentativo precedente potrebbero non rientrare: ogni "
            "nuovo tentativo richiede un nuovo run identifier.") in _rows(result)[0]["Pre conditions"]
    assert ("OnceAndDone è \"No re-entry\"" in report) is (mode == "OnceAndDone")
    result.package.journey.entry_mode = "MultipleEntries"
    assert "run identifier" not in _rows(result)[0]["Pre conditions"]


def test_entry_filter_goals_and_exits_are_reported():
    result = build_result()
    result.package.entry_event.entry_filter = "<FilterDefinition/>"
    result.package.journey.goal_count, result.package.journey.exit_count = 1, 2
    text = ("Il journey ha un filtro di ingresso che il generatore non modella: alcuni contatti generati potrebbero non "
            "entrare. Il journey ha 1 goal e 2 criteri di uscita che il generatore non modella: alcuni contatti "
            "potrebbero uscire prima della fine del percorso atteso. Se un contatto non segue il percorso atteso, "
            "controlla per primi questi elementi.")
    assert text in _how_to_run(write_report(result))
    assert all(text in row["Note"] for row in _rows(result))
    assert "goal" not in write_report(build_result())


def test_absent_rows_in_the_profile_and_the_preconditions():
    result = build_result()
    result.coverage.row_fields = {MASTER: (MASTER, "ContactId")}
    case_open, case_plain = [c for s in result.coverage.scenarios for c in s.cases if s.id == "S01"]
    case_open.contact_values[(MASTER, "ContactId")] = GENERATED_KEY
    case_plain.contact_values = {(MASTER, "Perm"): None, (MASTER, "ContactId"): None}
    first, second = _rows(result)[:2]
    assert "Profilo del contatto: DE_Master.Perm = True. Riga presente in DE_Master." in first["Description"]
    assert "Nessuna riga" not in first["Pre conditions"]
    assert "Profilo del contatto: nessuna riga in DE_Master." in second["Description"]
    assert ("Nessuna riga in DE_Master per questi contatti: DE_Master.csv non contiene le loro chiavi, perché il "
            "percorso atteso lo richiede.") in second["Pre conditions"]
    assert GENERATED_KEY not in first["Description"] + second["Description"]


def test_expected_result_past_an_unsupported_element():
    result = build_result()
    case_plain = next(c for c in result.coverage.cases if c.id == "S01-C02")
    case_plain.unknown_from = "WAIT-3"
    case_plain.unknown_reason = "l'attività 'WAIT-3' è di un tipo non supportato, SMSSYNC"
    second = _rows(result)[1]
    assert second["Test name"] == "S01-C02 - Percorso non noto da Wait one minute in poi"
    assert ("Percorso atteso: Channel [Email] > Welcome > Wait campaign date > Wait reminder date > Opened Welcome [No] "
            "> Permission [Has permission] > Wait one minute, da cui il percorso non è noto.") in second["Description"]
    expected = second["Expected result"]
    assert "Ultima attività raggiunta" not in expected
    assert "- Email ricevute prima di Wait one minute (WAIT-3): Welcome." in expected
    assert ("- Percorso non noto da Wait one minute (WAIT-3) in poi, perché l'attività 'WAIT-3' è di un tipo non "
            "supportato, SMSSYNC. Annota l'ultima attività raggiunta da ogni contatto.") in expected
    assert "| non nota da Wait one minute (WAIT-3) in poi |" in write_report(result)
    random_case = next(c for c in result.coverage.cases if c.random is not None)
    random_case.random.branches[1] = replace(random_case.random.branches[1], unknown_from="ENG-2",
                                             unknown_reason="lo split 'ENG-2' non è supportato")
    assert ("- 10%: email ricevute: Promo; valori scritti: nessuno; percorso non noto da Clicked Promo (ENG-2) in poi, "
            "perché lo split 'ENG-2' non è supportato.") in _rows(result)[2]["Expected result"]


def test_report_opens_with_a_warning_when_unsupported_elements_leave_branches_not_covered():
    assert "Attenzione" not in write_report(build_result())
    result = build_result()
    by_ref = {(b.activity_key, b.outcome_key): b for b in result.coverage.branches}
    for outcome, reason in (("perm", "branch non supportato: legge un campo senza relazione"),
                            ("remainder_path", "a valle di un elemento non supportato: questo branch viene dopo il "
                                               "branch non supportato 'Has permission' dello stesso split, quindi non "
                                               "viene mai valutato")):
        branch = by_ref[("DEC-2", outcome)]
        branch.status, branch.reason, branch.covered_by = "not covered", reason, []
    result.package.journey.activities["DEC-2"].outcome("perm").unsupported_reason = "legge un campo senza relazione"
    result.coverage.unsupported = ["DEC-2/perm"]
    report = write_report(result)
    head = report[:report.index("## Riepilogo")]
    assert ("> **Attenzione:** 2 branch su 10 non sono coperti a causa di un elemento non supportato: il branch 'Has "
            "permission' di 'Permission' (DEC-2). Per questi branch i dati di test non indicano il percorso dei "
            "contatti. Le sezioni Branch non coperti ed Elementi non supportati spiegano il motivo.") in head
    assert ("- Il branch 'Has permission' dello split Permission (DEC-2) non è supportato: legge un campo senza "
            "relazione. Questo branch, i branch che lo seguono nello split e quelli a valle non sono coperti.") in report
    assert unsupported_warning(build_result()) is None


def test_manual_action_windows_and_minimum_durations():
    result = build_result()
    case_open, case_plain = [c for s in result.coverage.scenarios for c in s.cases if s.id == "S01"]
    case_open.action_windows = [ActionWindow("ENG-1", "EMAIL-1", ["WAIT-X"], minutes=5, short=True)]
    case_open.min_duration_minutes = 5
    case_plain.min_duration_minutes, case_plain.not_before = 1, "2026-10-12 09:30"
    first, second = _rows(result)[:2]
    assert ("4. Entro 5 minuti dall'arrivo dell'email, apri l'email 'Welcome' inviata a ogni contatto di questo test"
            in first["Detailed steps"])
    assert "Dopo questo tempo i contatti raggiungono lo split 'Opened Welcome' (ENG-1)." in first["Detailed steps"]
    assert first["Note"] == ("Azione manuale richiesta: aprire l'email 'Welcome' entro 5 minuti dall'arrivo. La "
                             "finestra è sotto l'ora: apri l'email appena arriva.")
    assert ("Durata minima del test: 5 minuti; i contatti non finiscono il percorso prima del 2026-10-12 09:30."
            in first["Detailed steps"])                   # Case.not_before, filled by datagen
    assert ("Durata minima del test: 1 minuto; i contatti non finiscono il percorso prima del 2026-10-12 09:30."
            in second["Detailed steps"])
    assert "open di 'Welcome' entro 5 minuti dall'arrivo" in write_report(result)
    # A window that ends at a wait by attribute: the deadline datagen computed.
    case_open.action_windows = [ActionWindow("ENG-1", "EMAIL-1", ["WAIT-1", "WAIT-2"], minutes=0,
                                             attribute_wait_key="WAIT-2", deadline="2026-10-12 10:30")]
    assert "4. Prima del 2026-10-12 10:30, apri l'email 'Welcome'" in _rows(result)[0]["Detailed steps"]
    # A short window because a wait before the email delays it: the note says when the email arrives.
    case_open.action_windows = [ActionWindow("ENG-1", "EMAIL-1", ["WAIT-2"], minutes=0, attribute_wait_key="WAIT-2",
                                             deadline="2026-10-13 09:30", short=True, wait_keys_before=["WAIT-0"],
                                             email_not_before="2026-10-13 09:30")]
    note = _rows(result)[0]["Note"]
    assert "L'email arriva non prima del 2026-10-13 09:30, quindi la scadenza è a meno di un'ora dal suo arrivo" in note
    assert "dall'ora di generazione" not in note
    # A date already past holds nobody, and an unknown duration is said.
    case_plain.min_duration_minutes, case_plain.not_before = 0, "2026-10-11 09:30"
    assert ("Durata minima del test: nessuna, perché i wait del percorso possono terminare subito."
            in _rows(result)[1]["Detailed steps"])
    case_plain.min_duration_minutes = None
    assert "Durata minima del test non nota" in _rows(result)[1]["Detailed steps"]


def test_duration_of_a_path_unknown_past_an_activity_covers_only_the_known_part():
    """Spec, "Honest outcomes": past an unsupported element the path is unknown, so are its waits."""
    result = build_result()
    case_open, case_plain = [c for s in result.coverage.scenarios for c in s.cases if s.id == "S01"]
    for case in (case_open, case_plain):
        case.unknown_from = case.steps[-1].activity_key
    first, second = _rows(result)[:2]
    assert ("La durata vale solo per la parte nota del percorso, fino a Log open (UPD-1): dopo, il percorso non è "
            "noto e può avere altri wait.") in first["Detailed steps"]
    assert "La durata vale solo per la parte nota del percorso" in second["Detailed steps"]


def test_report_lists_only_the_assumptions_the_export_triggers():
    result = build_result()
    section = lambda: (lambda r: r[r.index("## Assunzioni"):r.index("## Impostazioni usate")])(write_report(result))  # noqa: E731
    before = section()
    assert "Un wait by attribute la cui data è già passata termina subito." in before
    assert "vince il primo" not in before                 # no decision split with two conditions
    assert "group-to-set" not in before and "uno a molti" not in before
    for activity in result.package.journey.activities.values():
        if activity.kind == KIND_WAIT_ATTRIBUTE:
            activity.kind = KIND_WAIT_DURATION
    result.package.relationships.append(Relationship("Contact -> Master", None, None, MASTER, "ContactId", "One", "Many"))
    after = section()
    assert "wait by attribute" not in after and "fuso orario" not in after
    assert ("La relazione group-to-set che collega DE_Master al contatto usa la Contact Key (leftRelationshipIDs 3)"
            in after)
    assert "Una condizione su una tabella in relazione uno a molti (DE_Master)" in after
    # An Update contact that writes into a table with the "row present / row absent" variable inserts its row.
    assert "riga mancante" not in after
    update = next(a for a in result.package.journey.activities.values() if a.kind == KIND_UPDATE)
    update.writes.append(Write(MASTER, "Perm", "True"))
    result.coverage.row_fields = {MASTER: (MASTER, "ContactId")}
    assert ("Un Update contact che inserisce la riga mancante del contatto in DE_Master dà ai campi che non scrive il "
            "loro valore di default, o null se non ne hanno.") in section()


def test_update_contact_whose_insert_may_fail_is_listed_as_unsupported_for_some_contacts():
    result = build_result()
    update = next(a for a in result.package.journey.activities.values() if a.kind == KIND_UPDATE)
    result.coverage.unsupported.append(update.key)
    result.coverage.unsupported_reasons[update.key] = "l'Update contact può fallire"
    report = write_report(result)
    assert ("non è supportata per una parte dei contatti: l'Update contact può fallire. Per questi contatti il "
            "percorso non è noto da quell'attività in poi") in report
    assert "che il generatore non supporta" not in report


# ---------------------------------------------------------------------------
# Spec, "Extensions for time-dependent journeys" (step 3): deadline, exit criteria, timeline, SMS, managed tables
# ---------------------------------------------------------------------------

ENTRY_TIME = datetime(2026, 10, 11, 12, 0)


def _case_row(result, case_id: str) -> dict:
    return next(row for row in _rows(result) if row["Test name"].startswith(case_id))


def test_report_opens_with_the_activation_deadline_and_the_testbook_gives_each_case_its_own():
    result = build_result()
    assert "attiva il journey entro" not in write_report(result)                # nothing depends on the clock
    assert outputs.activation_deadline_text(result) is None
    result.settings.entry_time = ENTRY_TIME
    result.coverage.validity_days = 2
    result.coverage.perishable = ["DE_Config.Start = 2026-10-13 12:00: resta valido 2 giorni dall'ora di ingresso"]
    day_level = ("il branch 'Email' dello split 'Channel' (DEC-1) è percorso dal caso S01-C02, in cui DE_Config.Start = "
                 "2026-10-13 12:00 cade esattamente sull'estremo della condizione X quando il contatto arriva allo "
                 "split: il branch è deciso sull'estremo di un confronto sul giorno")
    result.coverage.warnings.append(day_level)
    sentence = ("Importa i file e attiva il journey entro 2026-10-13 12:00: i valori scelti rispetto all'ora di ingresso "
                "restano validi 2 giorni.")
    report = write_report(result)
    assert outputs.activation_deadline_text(result) == sentence
    assert f"> **Attenzione:** {sentence}" in report[:report.index("## Riepilogo")]
    assert ("### Valori che dipendono dall'ora di ingresso\n\n- DE_Config.Start = 2026-10-13 12:00: resta valido 2 giorni "
            "dall'ora di ingresso.") in report
    assert "| Ora di ingresso | 2026-10-11 12:00 |" in report
    assert "| Validità minima dei dati (giorni) | 1 giorno |" in report
    for case in result.coverage.cases:
        case.validity_days = 1 if case.id == "S01-C02" else None
    first, second = _case_row(result, "S01-C01"), _case_row(result, "S01-C02")
    assert "Attiva entro" not in (first["Note"] or "")
    assert ("Attiva entro 2026-10-12 12:00: i valori di questo test restano validi 1 giorno dall'ora di ingresso "
            "(2026-10-11 12:00)") in second["Note"]
    assert "il branch è deciso sull'estremo di un confronto sul giorno" in second["Note"]   # only the case it names
    assert "assunzione" not in second["Note"]
    assert "confronto sul giorno" not in (first["Note"] or "")


def test_day_level_comparison_is_a_confirmed_fact_and_only_hours_and_minutes_stay_an_assumption():
    """Spec, "Relative dates in conditions": the day-level comparison is confirmed by the team (14 October 2026);
    the comparison of the full date and time with hours or minutes is still an assumption (docs/assumptions.md 48)."""
    from types import SimpleNamespace
    from cjhelper.model import RelativeDate

    result = SimpleNamespace(package=SimpleNamespace(journey=SimpleNamespace(exits=[])))
    day = [(None, None, SimpleNamespace(relative=RelativeDate(-1, 7, "days")))]
    hour = [(None, None, SimpleNamespace(relative=RelativeDate(-1, 1, "hours")))]
    assert outputs._time_assumptions(result, day) == []
    assert outputs._time_assumptions(result, day + hour) == [
        "Una condizione con una data relativa in ore o minuti confronta data e ora intere, non solo il giorno."]


def test_exit_criteria_section_summary_rows_and_what_is_still_not_modelled():
    result = build_result()
    assert "## Criteri di uscita" not in write_report(result)
    journey = result.package.journey
    journey.exits = [JourneyFilter("EXIT-1", "EXIT-1", "exit", None, description="Contract ended"),
                     JourneyFilter("EXIT-2", "EXIT-2", "exit", None, description="Odd rule",
                                   unsupported_reason="un operatore non supportato")]
    journey.exit_count = 2
    result.coverage.exit_targets = [BranchStatus("EXIT-1", "Contract ended", "exit", "met", "criteri di uscita soddisfatti",
                                                 STATUS_COVERED, "", ["S01-C02"])]
    report = write_report(result)
    assert report.index("## Branch non coperti") < report.index("## Criteri di uscita") < report.index("## Branch coperti con")
    section = report[report.index("## Criteri di uscita"):report.index("## Branch coperti con")]
    assert "| EXIT-1: Contract ended | valutato alla fine di ogni wait | coperto dal caso S01-C02 |" in section
    assert "| EXIT-2: Odd rule | non modellato: un operatore non supportato | non applicabile |" in section
    assert "| Target dei criteri di uscita | 1 |" in report and "| Target dei criteri di uscita coperti | 1 |" in report
    how = _how_to_run(report)
    assert "Il journey ha 1 criterio di uscita che il generatore non modella" in how      # only the second one
    assert "2 criteri di uscita" not in how


def test_testbook_writes_the_timeline_and_the_exit_by_criteria_last():
    result = build_result()
    result.settings.entry_time = ENTRY_TIME
    result.package.journey.exits = [JourneyFilter("EXIT-1", "EXIT-1", "exit", None, description="Contract ended")]
    assert "Giorno" not in _case_row(result, "S01-C02")["Detailed steps"]                    # a case without a timeline
    case = next(c for c in result.coverage.cases if c.id == "S01-C02")
    case.timeline = [TimelinePoint("DEC-1", "email", ENTRY_TIME), TimelinePoint("EMAIL-1", "e1", ENTRY_TIME),
                     TimelinePoint("WAIT-1", "w1", ENTRY_TIME + timedelta(days=1, hours=2)),
                     TimelinePoint("WAIT-3", None, ENTRY_TIME + timedelta(days=3))]
    case.exit_by_criteria = "WAIT-3"
    row = _case_row(result, "S01-C02")
    assert ("3. Controlla che il percorso rispetti questi tempi, contati dall'ora di ingresso (2026-10-11 12:00):\n"
            "   Giorno 0 (2026-10-11 12:00): 'Channel' (DEC-1), branch 'Email'; invio dell'email 'Welcome' (EMAIL-1)\n"
            "   Giorno 1 (2026-10-12 14:00): fine del wait 'Wait campaign date' (WAIT-1)\n"
            "   Giorno 3 (2026-10-14 12:00): il contatto esce dal journey alla fine del wait 'Wait one minute' (WAIT-3) "
            "per i criteri di uscita (Contract ended)\n4. ") in row["Detailed steps"]
    assert row["Expected result"].splitlines()[-1] == (
        "- Uscita: i contatti escono dal journey alla fine del wait 'Wait one minute' (WAIT-3), perché soddisfano i "
        "criteri di uscita (Contract ended).")


def test_sms_step_and_result_name_the_contact_with_the_real_number_or_say_it_cannot_be_checked():
    result = build_result()
    result.package.journey.activities["SMS-1"] = Activity(key="SMS-1", raw_type="SMSSYNC", kind=KIND_SMS,
                                                          name="Reminder SMS", outcomes=[Outcome("s1", "", None)])
    result.package.journey.default_mobile_field = "Phone"
    case = next(c for c in result.coverage.cases if c.id == "S01-C02")
    case.sms_sent = ["SMS-1"]
    row = _case_row(result, "S01-C02")
    assert "non verificabile senza numero reale" in row["Expected result"]
    assert "l'SMS non può essere verificato" in row["Detailed steps"]
    assert "Il caso S01-C02 invia un SMS, ma non è stato indicato un numero di telefono reale: l'invio dell'SMS non " \
           "può essere verificato." in write_report(result)
    result.settings.real_phone = "393401234567"
    for contact in result.data.contacts:
        contact.uses_real_phone = contact.key == case.contact_keys[0]
    row = _case_row(result, "S01-C02")
    assert ("SMS 'Reminder SMS' (SMS-1) inviato al numero del contatto TST_R01_C0004, che ha il numero reale "
            "393401234567 nella colonna Phone di DE_Entry.") in row["Detailed steps"]
    assert ("- SMS ricevuti: 'Reminder SMS' (SMS-1) sul telefono del contatto TST_R01_C0004") in row["Expected result"]
    assert "SMS" not in _case_row(result, "S01-C01")["Detailed steps"]


def test_managed_table_is_imported_like_the_others_and_noted_as_synchronized():
    """Spec, "Synchronized data extensions" (confirmed by the team on 14 October 2026): the CSV is in the import
    list; the report only notes that the table is synchronized from the CRM."""
    result = build_result()
    table = next(t for t in result.package.tables.values() if t.name == MASTER)
    unmarked = write_report(result)
    table.managed = True
    assert not any("sincronizzat" in item for item in outputs._triggered_assumptions(result))   # no assumption left
    report = write_report(result)
    how_to_run = _how_to_run(report)
    assert "Importa prima DE_Master.csv e DE_Config.csv, poi DE_Entry.csv" in how_to_run                 # in the list
    assert ("DE_Master è una data extension sincronizzata dal CRM: il CSV si importa come gli altri e le sue righe "
            "arriveranno anche dalla sincronizzazione.") in how_to_run
    assert "| DE_Master | source, split | contatto | DE_Master.csv |" in report
    assert ("DE_Master è una data extension sincronizzata dal CRM (managed data extension, che Package Manager "
            "segnala come NO_DEPLOY o ManagedDataExtension): il CSV si importa come gli altri e le sue righe "
            "arriveranno anche dalla sincronizzazione.") in report                                    # the Tabelle section
    pre = _rows(result)[0]["Pre conditions"]
    assert "Importa questi file CSV: DE_Entry.csv, DE_Master.csv, DE_Config.csv. Importa DE_Entry.csv per ultimo." in pre
    assert "DE_Master è una data extension sincronizzata dal CRM: il CSV si importa" in pre
    assert "DE_Master.csv" in result.files and result.data.tables[MASTER].rows
    assert not any(fragment in report + pre for fragment in ("non importabile", "Non importare"))
    assert "sincronizzata dal CRM" not in unmarked
    table.managed = False
    assert "sincronizzat" not in write_report(result)


def test_link_role_is_written_in_italian_and_isnot_is_a_text_operator():
    assert outputs.role_text({ROLE_LINK, "source"}) == "collegamento, source"
    assert outputs.role_text(set()) == "nessuno"
    result = build_result()
    result.table_model.tables[MASTER].roles.add(ROLE_LINK)
    report = write_report(result)
    assert "| DE_Master | collegamento, source, split |" in report
    assert "Una tabella di collegamento sta sul percorso" in report
    assert "isnot" in outputs._CASE_OPERATORS

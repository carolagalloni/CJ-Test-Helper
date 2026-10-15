"""Write the generated files: one CSV per table, scenarios.csv, report.md, testbook.xlsx and the zip.

Spec, "Data generation rules", file format: "UTF-8, comma-separated, header row with the exact field
names in schema order, True/False for Booleans, empty cell for null."

Spec, "Outputs": "The app returns one zip file with the data CSVs plus three supporting files":
``<TableName>.csv`` (one per selected table), ``scenarios.csv``, ``testbook.xlsx`` and ``report.md``.

Every function here is pure: the same input gives the same bytes ("The output is reproducible").
For that reason the timestamps that Excel files and zip archives normally carry are fixed.

Spec, "Extensions beyond the pilot" (12 October 2026), as report.md and the testbook apply it:

- Entry source and procedure: the procedure follows ``EntryEvent.kind``. Automation (the pilot): as before.
  Run Once: import every CSV, then activate the journey; only rows present at activation enter; a new attempt
  needs a new journey version and a new run identifier; all scenarios enter together; the entry data extension
  must hold only test rows when the journey is activated. Recurring schedule: a procedure to verify. CloudPages
  form, API event or any other: a copy of the journey with a data extension entry with the same columns.
  ``OnceAndDone`` and ``SingleEntryAcrossAllVersions``: every attempt needs a new run identifier. Entry filters,
  goals and exit criteria are not modelled and are reported.
- Row present or absent: a contact whose case has the row absent is described as having no row in that table.
- Honest outcomes: past an unsupported element the expected result says the path is unknown from that activity
  on, and never names it as the last activity reached; report.md (and the page, through
  ``unsupported_warning``) opens with a warning when unsupported elements leave branches not covered; the
  deadline of a manual action comes from the waits between the email and the engagement split
  (``Case.action_windows``); each test gives its minimum duration; report.md lists only the assumptions the
  export triggers.

Spec, "Extensions for time-dependent journeys" (14 October 2026), as report.md and the testbook apply it:

- Validity of the data: report.md opens with the deadline for importing and activating (the entry time plus
  ``CoverageResult.validity_days``) and lists the perishable values; each test that has a ``Case.validity_days``
  gives its own deadline in the Note.
- The clock: the steps of a test list the timeline of its case, "Giorno N (date and time): ...".
- Exit criteria: a section of report.md says which are evaluated at the end of every wait, which are not
  modelled, and the state of their coverage target; a case that leaves the journey that way says so last.
- SMS sends: the steps and the expected result name the SMS and the contact that receives it (the one with the
  real number), or say it cannot be checked without a real number.
- Synchronized data extensions (managed): their CSV is imported like the others; a note says the table is
  synchronized from the CRM.
"""

from __future__ import annotations

import csv
import dataclasses
import io
import math
import re
import zipfile
from datetime import datetime, timedelta, timezone
from typing import Iterable, Optional

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from . import criteria
from .datagen import (
    SCENARIO_COLUMNS,
    UNDELIVERABLE_DOMAIN,
    activation_setting,
    boolean_text,
    one_active_not_offered,
    case_paths,
    wait_precedes_manual_action,
)
from .graph import is_split
from .model import (
    ENTRY_AUTOMATION,
    ENTRY_OTHER,
    ENTRY_RUN_ONCE,
    ENTRY_SCHEDULED,
    KIND_DECISION,
    KIND_EMAIL,
    KIND_ENGAGEMENT,
    KIND_RANDOM,
    KIND_SMS,
    KIND_UNSUPPORTED,
    KIND_UPDATE,
    KIND_WAIT_ATTRIBUTE,
    KIND_WAIT_DURATION,
    KIND_WAIT_STO,
    LEVEL_SCENARIO,
    ROLE_LINK,
    STATUS_COVERED,
    STATUS_MANUAL,
    STATUS_NOT_COVERED,
    Activity,
    ActionWindow,
    BranchStatus,
    Case,
    EqualityFilter,
    FieldRef,
    GenerationResult,
    Journey,
    Outcome,
    Package,
    Scenario,
    Settings,
    Step,
    Table,
    TableData,
    UnsupportedFilter,
    Value,
    Write,
    formula_conditions,
    formula_text,
)

CSV_LINE_END = "\n"
SCENARIOS_FILE = "scenarios.csv"
REPORT_FILE = "report.md"
TESTBOOK_FILE = "testbook.xlsx"

# Spec, "Columns of testbook.xlsx".
TESTBOOK_COLUMNS = ["ID", "Owner system", "Test name", "Description", "Pre conditions",
                    "Detailed steps", "Expected result", "Status", "Note"]
OWNER_SYSTEM = "SFMC"
STATUS_TO_START = "TO START"
TESTBOOK_SHEET = "Testbook"
_TESTBOOK_WIDTHS = [6, 12, 40, 60, 60, 70, 70, 12, 40]

ZIP_DATE_TIME = (1980, 1, 1, 0, 0, 0)    # fixed, so the same files always give the same zip bytes
_CORE_TIMESTAMP = re.compile(r"(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)")
_CURRENT_TIME_VALUE = "{{DateTime.Now}}"
_REPORT_DETAIL_LIMIT = 50
SUSPICIOUS_PREFIX = "pattern sospetto:"    # how other modules mark a suspicious pattern


def table_file_name(table_name: str) -> str:
    """Spec, "Outputs": one ``<TableName>.csv`` per selected table."""
    return f"{table_name}.csv"


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------

def schema_field_names(table: Table) -> list[str]:
    """The exact field names in schema order."""
    return [f.name for f in sorted(table.fields, key=lambda f: f.ordinal)]


def _csv_bytes(header: list[str], rows: Iterable[list[Value]]) -> bytes:
    """UTF-8 without byte order mark, comma separated, minimal quoting, "\\n" line ends, None -> empty cell."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator=CSV_LINE_END)
    writer.writerow(header)
    for values in rows:
        writer.writerow(["" if value is None else value for value in values])
    return buffer.getvalue().encode("utf-8")


def _row_value(row: dict[str, Value], column: str) -> Value:
    if column in row:
        return row[column]
    low = column.lower()
    for key, value in row.items():
        if key.lower() == low:
            return value
    return None


def write_table_csv(data: TableData, table: Table) -> bytes:
    """One data CSV: the header is the table's schema, in schema order, whatever the rows hold."""
    header = schema_field_names(table)
    return _csv_bytes(header, ([_row_value(row, column) for column in header] for row in data.rows))


def write_scenarios_csv(rows: list[dict[str, Value]]) -> bytes:
    """scenarios.csv: one row per scenario; the columns are those of the rows, in first-seen order."""
    header: list[str] = []
    for row in rows:
        header.extend(column for column in row if column not in header)
    if not header:
        header = list(SCENARIO_COLUMNS)
    return _csv_bytes(header, ([row.get(column) for column in header] for row in rows))


def read_csv(data: bytes) -> tuple[list[str], list[dict[str, Value]]]:
    """Read a CSV written here back: (header, rows). An empty cell reads as None.

    Strict UTF-8, so a byte order mark would show up in the first column name. Blank lines are
    skipped, cells beyond the header are dropped and missing cells read as None.
    """
    reader = csv.reader(io.StringIO(data.decode("utf-8"), newline=""))
    lines = list(reader)
    if not lines:
        return [], []
    header = lines[0]
    rows = [{column: (line[i] if i < len(line) and line[i] != "" else None) for i, column in enumerate(header)}
            for line in lines[1:] if line]
    return header, rows


# ---------------------------------------------------------------------------
# Zip
# ---------------------------------------------------------------------------

def _zip_info(name: str) -> zipfile.ZipInfo:
    info = zipfile.ZipInfo(name, date_time=ZIP_DATE_TIME)
    info.compress_type = zipfile.ZIP_DEFLATED
    info.external_attr = 0o644 << 16
    return info


def build_zip(files: dict[str, bytes]) -> bytes:
    """Spec, "Outputs": "one zip file with the data CSVs plus three supporting files". Files keep the
    given order; every entry has the same fixed timestamp."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(_zip_info(name), content)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Shared wording helpers
# ---------------------------------------------------------------------------

_NAME_LIMIT = 90         # a test name stays within this length (spec example: "S07-C02 - Email sent on channel 1, ...")
_LABEL_LIMIT = 24        # an outcome label longer than this is shortened inside a test name
_FORMULA_LIMIT = 120     # a condition longer than this is not quoted in a reason
_PLAIN_FIRST_WORD = re.compile(r"[a-z]+(?=[\s,:;]|$)")

# Activity kinds in words, for the expected result.
_KIND_WORDS = {
    KIND_DECISION: "questo decision split",
    KIND_ENGAGEMENT: "questo engagement split",
    KIND_RANDOM: "questo random split",
    KIND_WAIT_DURATION: "questo wait",
    KIND_WAIT_ATTRIBUTE: "questo wait",
    KIND_WAIT_STO: "questo send-time wait",
    KIND_EMAIL: "questa email",
    KIND_UPDATE: "questo Update contact",
}

# Duration units as the export writes them (singular, lower case) -> Italian singular and plural.
_UNIT_WORDS = {
    "second": ("secondo", "secondi"),
    "minute": ("minuto", "minuti"),
    "hour": ("ora", "ore"),
    "day": ("giorno", "giorni"),
    "week": ("settimana", "settimane"),
    "month": ("mese", "mesi"),
    "year": ("anno", "anni"),
}

# A manual action (``ManualAction.kind``) in the forms the testbook needs.
_ACTION_INFINITIVE = {"open": "aprire", "click": "cliccare"}
_ACTION_TO_DO = {"open": "aprire l'email", "click": "cliccare un link nell'email"}
_ACTION_IMPERATIVE = {"open": "apri l'email", "click": "clicca un link nell'email"}
_ACTION_SUBJUNCTIVE = {"open": "aprano l'email", "click": "clicchino un link nell'email"}

# The level of a table as the report writes it.
_LEVEL_WORDS = {"contact": "contatto", "scenario": "scenario"}

# The not-covered states of ``_explain_reason`` after "che" ("which is ...").
_STATE_WORDS = {"unreachable": "non è raggiungibile", "not covered": "non è coperto"}


def _show(value: Value) -> str:
    return "(vuoto)" if value is None or value == "" else str(value)


def _plural(count: int, singular: str, plural: str) -> str:
    return f"{count} {singular if count == 1 else plural}"


def _covered(count: int, words: str = "coperto") -> str:
    """"1 è coperto", "8 sono coperti", "1 non è coperto"; ``words`` is the singular adjective, with a
    leading "non " when negated."""
    negated = words.startswith("non ")
    adjective = words[4:] if negated else words
    verb = "è" if count == 1 else "sono"
    shown = adjective if count == 1 else adjective[:-1] + "i"
    return f"{count} {'non ' if negated else ''}{verb} {shown}"


def _lie_behind(count: int) -> str:
    """"1 branch si trova", "3 branch si trovano"."""
    return f"{count} branch {'si trova' if count == 1 else 'si trovano'}"


def _join_words(items: Iterable[str]) -> str:
    """'a', 'a e b', 'a, b e c'."""
    items = list(items)
    if len(items) <= 1:
        return "".join(items)
    return ", ".join(items[:-1]) + " e " + items[-1]


def _sentence(text: str) -> str:
    """A message as a sentence: a plain first word gets a capital letter and the text ends with a full
    stop. A name or key at the start is left as it is."""
    text = str(text).strip()
    if not text:
        return text
    if _PLAIN_FIRST_WORD.match(text):
        text = text[0].upper() + text[1:]
    return text if text.endswith((".", "!", "?")) else text + "."


def _table_of(package: Package, table_name: str) -> Optional[Table]:
    return next((t for t in package.tables.values() if t.name.lower() == table_name.lower()), None)


def _field_type(package: Package, table_name: str, field_name: str) -> str:
    table = _table_of(package, table_name)
    field = table.field(field_name) if table is not None else None
    return field.type.lower() if field is not None and field.type else ""


def _show_value(package: Package, table_name: str, field_name: str, value: Value) -> str:
    """A value as the CSV holds it: Booleans as True/False."""
    if _field_type(package, table_name, field_name) == "boolean":
        value = boolean_text(value) or value
    return _show(value)


def _ordered_refs(package: Package, refs: Iterable[FieldRef]) -> list[FieldRef]:
    """Fields grouped by table, tables in first-seen order, fields in schema order: the CSV column order."""
    refs = list(refs)
    tables: list[str] = []
    for table, _ in refs:
        if table.lower() not in tables:
            tables.append(table.lower())

    def position(ref: FieldRef) -> tuple[int, int]:
        table = _table_of(package, ref[0])
        field = table.field(ref[1]) if table is not None else None
        return tables.index(ref[0].lower()), (field.ordinal if field is not None else 10 ** 6)

    return sorted(refs, key=position)


def _values_parts(package: Package, values: dict, order: list) -> tuple[Optional[str], list[str]]:
    """``Field = value`` items in CSV column order. When the fields are several and all in one table,
    that table is returned and named once by the caller; otherwise each item reads ``Table.Field = value``."""
    refs = [ref for ref in order if ref in values] + [ref for ref in values if ref not in order]
    refs = _ordered_refs(package, refs)
    one_table = refs[0][0] if len(refs) > 1 and len({table.lower() for table, _ in refs}) == 1 else None
    items = []
    for table, field in refs:
        name = field if one_table else f"{table}.{field}"
        items.append(f"{name} = {_show_value(package, table, field, values[(table, field)])}")
    return one_table, items


def _duration_text(amount, unit: Optional[str], english: bool = False) -> str:
    unit_text = (unit or "").lower()
    if amount == 1 and unit_text.endswith("s"):
        unit_text = unit_text[:-1]
    words = None if english else _UNIT_WORDS.get(unit_text[:-1] if unit_text.endswith("s") else unit_text)
    if words is not None:
        unit_text = words[0] if amount == 1 else words[1]
    return f"{amount if amount is not None else '?'} {unit_text}".strip()


def _date(moment: datetime, settings: Settings) -> str:
    return moment.strftime(settings.date_format)


def _human_date_format(date_format: str) -> str:
    """``%Y-%m-%d %H:%M`` as the spec writes it: ``YYYY-MM-DD HH:MM``. Other formats are shown as they are."""
    text = date_format
    for code, shown in (("%Y", "YYYY"), ("%m", "MM"), ("%d", "DD"), ("%H", "HH"), ("%M", "MM"), ("%S", "SS")):
        text = text.replace(code, shown)
    return date_format if "%" in text else text


def _activity_name(journey: Journey, key: Optional[str]) -> str:
    """An activity by its name, or its key when it has none."""
    if key is None:
        return "(nessuna)"
    activity = journey.activities.get(key)
    return activity.name if activity is not None and activity.name else key


def _activity_label(journey: Journey, key: Optional[str]) -> str:
    """``Name (KEY)``, or the key alone when the activity has no name."""
    name = _activity_name(journey, key)
    return name if name == key else f"{name} ({key})"


def _quoted_label(journey: Journey, key: Optional[str]) -> str:
    """An activity named inside a sentence: ``'Name' (KEY)``, or the key alone when it has no name."""
    name = _activity_name(journey, key)
    return name if name == key else f"'{name}' ({key})"


def _step_text(journey: Journey, step: Step) -> str:
    """An activity of the path; for an activity with several outcomes, the outcome taken."""
    text = _activity_name(journey, step.activity_key)
    activity = journey.activities.get(step.activity_key)
    if activity is not None and len(activity.outcomes) > 1 and step.outcome_key is not None:
        outcome = activity.outcome(step.outcome_key)
        text += f" [{outcome.label if outcome is not None and outcome.label else step.outcome_key}]"
    return text


def _path_text(journey: Journey, steps: Iterable[Step]) -> str:
    texts = [_step_text(journey, step) for step in steps]
    return " > ".join(texts) if texts else "(nessuna attività)"


def _times(count: int) -> str:
    return "una volta" if count == 1 else f"{count} volte"


def _emails_text(journey: Journey, keys: Iterable[str]) -> str:
    """The emails received, by name, with a count when one is received more than once: "Demo Sky, 2 volte"."""
    counts: dict[str, int] = {}
    for key in keys:
        activity = journey.activities.get(key)
        name = (activity.email_name or activity.name or key) if activity is not None else key
        counts[name] = counts.get(name, 0) + 1
    if not counts:
        return "nessuna"
    if len(counts) == 1:
        name, count = next(iter(counts.items()))
        return name if count == 1 else f"{name}, {count} volte"
    return _join_words(f"{name} {_times(count)}" for name, count in counts.items())


def _current_time_fields(journey: Journey, keys: Iterable[str]) -> set[tuple[str, str]]:
    """Fields that an update contact step among these activities sets to ``{{DateTime.Now}}``."""
    fields: set[tuple[str, str]] = set()
    for key in keys:
        activity = journey.activities.get(key)
        for write in (activity.writes if activity is not None else []):
            if write.value == _CURRENT_TIME_VALUE:
                fields.add((write.table.lower(), write.field.lower()))
    return fields


def _writes_text(writes: Iterable[Write], current_time: set[tuple[str, str]] = frozenset(), now_text: str = "") -> str:
    """The final value of each field written by the journey, grouped by table in the order first written.
    A field set to ``{{DateTime.Now}}`` gets the time its step runs, not the generation time."""
    final: dict[tuple[str, str], Value] = {}
    for write in writes:
        final[(write.table, write.field)] = write.value
    groups: dict[str, list[str]] = {}
    for (table, field), value in final.items():
        if value == _CURRENT_TIME_VALUE or (now_text and value == now_text
                                            and (table.lower(), field.lower()) in current_time):
            shown = "la data e l'ora in cui viene eseguito il passo"
        else:
            shown = _show(value)
        groups.setdefault(table, []).append(f"{field} = {shown}")
    parts = [f"{table}.{items[0]}" if len(items) == 1 else f"{table}: {', '.join(items)}"
             for table, items in groups.items()]
    return "; ".join(parts) if parts else "nessuno"


def _ordered_scenarios(result: GenerationResult) -> list[Scenario]:
    return sorted(result.coverage.scenarios, key=lambda s: (s.run_order, s.id))


def _key_range(keys: list[str]) -> str:
    if not keys:
        return "nessuno"
    if len(keys) <= 3:
        return ", ".join(keys)
    return f"da {keys[0]} a {keys[-1]}"


def _smallest_branch(case: Case):
    branches = case.random.branches if case.random is not None else []
    return min(branches, key=lambda b: b.percentage) if branches else None


def _outcome_label(journey: Journey, activity_key: str, outcome_key: str) -> str:
    activity = journey.activities.get(activity_key)
    outcome = activity.outcome(outcome_key) if activity is not None else None
    return outcome.label if outcome is not None and outcome.label else outcome_key


def _branch_label(journey: Journey, random_key: str, branch) -> str:
    """A random split branch by its label, with its percentage when the label does not show it."""
    label = _outcome_label(journey, random_key, branch.outcome_key)
    return label if f"{branch.percentage}%" in label else f"{label} ({branch.percentage}%)"


def _range_chance(contacts: int, percentage: Optional[int], low: int, high: int) -> Optional[int]:
    """In how many runs out of 100 the smallest side of a random split gets low to high contacts, when each
    contact takes it with the split's percentage (spec: 1 to 9 of 50 on the 10% side, "about 97% of runs")."""
    if contacts <= 0 or not percentage or not 0 < percentage < 100 or high < low:
        return None
    p = percentage / 100
    chance = sum(math.comb(contacts, k) * p ** k * (1 - p) ** (contacts - k)
                 for k in range(max(low, 0), min(high, contacts) + 1))
    return round(chance * 100)


def _path_keys(case: Case) -> list[str]:
    """Every activity a case can cross, in path order: the deterministic part, then the random branches."""
    keys: list[str] = []
    for path in case_paths(case):
        for key in path:
            if key not in keys:
                keys.append(key)
    return keys


def _end_sentence(journey: Journey, last: Optional[Step]) -> str:
    """The end of a path in words, when the last activity leads out of the journey."""
    activity = journey.activities.get(last.activity_key) if last is not None else None
    if activity is None or activity.kind == KIND_UNSUPPORTED:
        return ""
    if last.outcome_key is not None:
        outcome = activity.outcome(last.outcome_key)
    else:
        outcome = activity.outcomes[0] if len(activity.outcomes) == 1 else None
    if outcome is None or outcome.next is not None:
        return ""
    if len(activity.outcomes) > 1:
        return f"Lì i contatti escono dal journey, sul branch {outcome.label or outcome.key}."
    return f"I contatti escono dal journey dopo {_KIND_WORDS.get(activity.kind, 'questa attività')}."


# ---------------------------------------------------------------------------
# Entry source, unknown paths, windows and durations (spec, "Extensions beyond the pilot")
# ---------------------------------------------------------------------------

# Entry modes with which a contact key of an earlier attempt may not enter again: "Entry modes OnceAndDone (no
# re-entry) and SingleEntryAcrossAllVersions: the testbook says every attempt needs a new run identifier".
SINGLE_ENTRY_MODES = ("OnceAndDone", "SingleEntryAcrossAllVersions")

# The not-covered reasons of coverage.py that come from an unsupported element (the branch itself, a branch
# after it in the same split, or a branch behind it).
_UNSUPPORTED_REASONS = ("attività non supportata", "branch non supportato", "a valle di un elemento non supportato")

_DAY_MINUTES = 24 * 60     # minutes in a day, for durations in words


def entry_kind(package) -> str:
    """The entry source kind (``EntryEvent.kind``, one of ``ENTRY_*``); without an entry event, the pilot's."""
    event = getattr(package, "entry_event", None)
    return getattr(event, "kind", None) or ENTRY_AUTOMATION


def entry_source_text(package) -> str:
    """An entry source that is not a scheduled data extension (``ENTRY_OTHER``), in words, with its type as
    the export writes it: "un form CloudPages (CloudPagesSmartCaptureSubmissionEvent)"."""
    event = getattr(package, "entry_event", None)
    event_type = (getattr(event, "event_type", "") or "").strip()
    low = event_type.lower()
    if low.startswith("cloudpages"):
        return f"un form CloudPages ({event_type})"
    if low.startswith("api"):
        return f"un API event ({event_type})"
    return f"un evento di tipo {event_type}" if event_type else "un evento di tipo non indicato"


def _entry_mode(package) -> Optional[str]:
    """The journey's entry mode when it keeps a contact of an earlier attempt from entering again."""
    mode = getattr(package.journey, "entry_mode", None)
    return mode if mode in SINGLE_ENTRY_MODES else None


def _not_modelled_text(package) -> Optional[str]:
    """Spec: "Entry filters, goals and exit criteria are not modelled; when present they are reported, because
    contacts may not enter or may leave early"."""
    journey = package.journey
    event = getattr(package, "entry_event", None)
    parts = []
    if getattr(event, "entry_filter", None):
        parts.append("Il journey ha un filtro di ingresso che il generatore non modella: alcuni contatti generati "
                     "potrebbero non entrare.")
    goals, exits = getattr(journey, "goal_count", 0), getattr(journey, "exit_count", 0)
    parsed = getattr(journey, "exits", None) or []
    if parsed:
        # Spec, step 3: the exit criteria the generator can evaluate are modelled; only the others are reported here.
        exits = sum(1 for exit_ in parsed if exit_.unsupported_reason is not None)
    if goals or exits:
        found = ([_plural(goals, "goal", "goal")] if goals else []) + \
                ([_plural(exits, "criterio di uscita", "criteri di uscita")] if exits else [])
        parts.append(f"Il journey ha {_join_words(found)} che il generatore non modella: alcuni contatti potrebbero "
                     "uscire prima della fine del percorso atteso.")
    if not parts:
        return None
    return " ".join(parts) + " Se un contatto non segue il percorso atteso, controlla per primi questi elementi."


def _unsupported_elements(result: GenerationResult) -> list[str]:
    """The unsupported activities and branches, by name: an unsupported branch is written ``<split>/<outcome>``
    in ``CoverageResult.unsupported``."""
    journey = result.package.journey
    labels = []
    for key in _unique(list(result.package.unsupported_activities) + list(result.coverage.unsupported)):
        activity_key, _, outcome_key = key.partition("/")
        if key not in journey.activities and outcome_key and activity_key in journey.activities:
            labels.append(f"il branch '{_outcome_label(journey, activity_key, outcome_key)}' di "
                          f"{_quoted_label(journey, activity_key)}")
        else:
            labels.append(_quoted_label(journey, key))
    return labels


def unsupported_warning(result: GenerationResult) -> Optional[str]:
    """Spec, "Honest outcomes": "report.md and the page open with a warning when unsupported elements leave
    branches not covered". None when no branch is left not covered by an unsupported element."""
    branches = list(result.coverage.branches)
    hit = [b for b in branches if b.status == STATUS_NOT_COVERED
           and (getattr(b, "reason", "") or "").startswith(_UNSUPPORTED_REASONS)]
    if not hit:
        return None
    elements = _unsupported_elements(result)
    shown = elements[:3] + ([f"altri {len(elements) - 3}"] if len(elements) > 3 else [])
    verb = "non è coperto" if len(hit) == 1 else "non sono coperti"
    cause = "un elemento non supportato" if len(elements) == 1 else "elementi non supportati"
    text = f"{len(hit)} branch su {len(branches)} {verb} a causa di {cause}"
    if shown:
        text += f": {_join_words(shown)}"
    return text + ". Per questi branch i dati di test non indicano il percorso dei contatti."


# ---------------------------------------------------------------------------
# Time-dependent journeys (spec, "Extensions for time-dependent journeys", 14 October 2026)
# ---------------------------------------------------------------------------

_ROLE_WORDS = {ROLE_LINK: "collegamento"}    # the report's role names are Salesforce's own words, except this one
_DAY_LEVEL_WORDS = "confronto sul giorno"    # how coverage.py names the day-level comparison in its bound notes
_EXIT_WORDS = "criteri di uscita"


def role_text(roles: Iterable[str]) -> str:
    """The roles of a table as report.md and the page write them: ``{"split", "source"}`` -> "source, split",
    and the role ``link`` as "collegamento"."""
    return ", ".join(sorted(_ROLE_WORDS.get(role, role) for role in roles)) or "nessuno"


def entry_time(result: GenerationResult) -> datetime:
    """Spec, "The clock": the clock starts at the entry time, a setting whose default is the generation time."""
    return result.settings.entry_time or result.generation_time


def activation_deadline(result: GenerationResult) -> Optional[datetime]:
    """Spec, "Validity of the data": the entry time plus the earliest validity over the cases; None when nothing
    depends on the clock."""
    days = getattr(result.coverage, "validity_days", None)
    return None if days is None else entry_time(result) + timedelta(days=days)


def activation_deadline_text(result: GenerationResult) -> Optional[str]:
    """The sentence report.md opens with and the page shows after the generation: by when to import the files
    and activate the journey. None when the data is not perishable."""
    deadline = activation_deadline(result)
    if deadline is None:
        return None
    days = result.coverage.validity_days
    valid = ("non tollerano ritardi di attivazione" if days <= 0
             else f"restano validi {_plural(days, 'giorno', 'giorni')}")
    return (f"Importa i file e attiva il journey entro {_date(deadline, result.settings)}: i valori scelti rispetto "
            f"all'ora di ingresso {valid}.")


def _one_line(text) -> str:
    return " ".join(str(text or "").split())


def exit_description(exit_) -> str:
    """An exit criterion in words: the description the export gives it, or its formula."""
    text = _one_line(exit_.description)
    if not text and exit_.formula is not None:
        text = formula_text(exit_.formula)
    return text or _one_line(exit_.name) or exit_.key


def exit_evaluation_text(exit_) -> str:
    """Whether the generator evaluates an exit criterion (at the end of every wait) or not, and why."""
    if exit_.unsupported_reason is None:
        return "valutato alla fine di ogni wait"
    return f"non modellato: {_one_line(exit_.unsupported_reason).rstrip('.')}"


def _exit_phrase(journey: Journey) -> str:
    """" (description; description)" of the exit criteria the generator evaluates, or nothing."""
    found = [exit_description(e) for e in getattr(journey, "exits", None) or [] if e.unsupported_reason is None]
    return f" ({'; '.join(found)})" if found else ""


def _managed_tables(result: GenerationResult) -> list[str]:
    """The tables with a CSV that Package Manager marks as managed (synchronized) data extensions: spec,
    "Synchronized data extensions", their rows also arrive from the CRM sync."""
    found = []
    for name in result.data.tables:
        table = _table_of(result.package, name)
        if table is not None and getattr(table, "managed", False):
            found.append(name)
    return found


def _managed_text(result: GenerationResult, names: list[str], short: bool = False) -> str:
    """The note on a managed (synchronized) table (spec, "Synchronized data extensions"): its CSV is imported like
    the others and its rows will also arrive from the sync; ``short`` is the form for the import instructions."""
    one = len(names) == 1
    who = _join_words(names)
    synced = "è una data extension sincronizzata dal CRM" if one else "sono data extension sincronizzate dal CRM"
    imported = "il CSV si importa" if one else "i CSV si importano"
    arrive = "le sue righe arriveranno" if one else "le loro righe arriveranno"
    if short:
        return f"{who} {synced}: {imported} come gli altri e {arrive} anche dalla sincronizzazione."
    return (f"{who} {synced} (managed data extension, che Package Manager segnala come NO_DEPLOY o "
            f"ManagedDataExtension): {imported} come gli altri e {arrive} anche dalla sincronizzazione.")


def _known_steps(steps: Iterable[Step], unknown_from: Optional[str]) -> list[Step]:
    """The steps of a path before the activity it stopped at without exiting (``Case.unknown_from``)."""
    steps = list(steps)
    if (unknown_from is not None and steps and steps[-1].activity_key == unknown_from
            and steps[-1].outcome_key is None):
        return steps[:-1]
    return steps


def _expected_path_text(journey: Journey, steps: Iterable[Step], unknown_from: Optional[str]) -> str:
    """The expected path; past the activity where it stops being known, the path is said to be unknown."""
    if unknown_from is None:
        return _path_text(journey, steps)
    known = _known_steps(steps, unknown_from)
    tail = f"{_activity_name(journey, unknown_from)}, da cui il percorso non è noto"
    return f"{_path_text(journey, known)} > {tail}" if known else tail


def _unknown_text(journey: Journey, key: str, reason: str) -> str:
    """Spec: "Past an unsupported element the expected result says the path is unknown from that activity on"."""
    text = f"Percorso non noto da {_activity_label(journey, key)} in poi"
    reason = (reason or "").strip().rstrip(".")
    return f"{text}, perché {reason}" if reason else text


def _minutes_words(minutes: int) -> str:
    """A duration in words: "4 giorni", "1 giorno e 5 minuti", "2 ore"."""
    days, rest = divmod(minutes, _DAY_MINUTES)
    hours, mins = divmod(rest, 60)
    parts = ([_plural(days, "giorno", "giorni")] if days else []) + ([_plural(hours, "ora", "ore")] if hours else [])
    if mins or not parts:
        parts.append(_plural(mins, "minuto", "minuti"))
    return _join_words(parts)


def _duration_sentence(case: Case, now: datetime, date_format: str) -> str:
    """Spec, "Honest outcomes": "The testbook gives each case's minimum duration": the fixed waits of the case's
    longest path (``Case.min_duration_minutes``) and the date before which it cannot end (``Case.not_before``),
    when that date is after the generation time."""
    minutes = getattr(case, "min_duration_minutes", 0)
    not_before = getattr(case, "not_before", None)
    try:
        if not_before and datetime.strptime(not_before, date_format) <= now:
            not_before = None      # a date already past holds nobody
    except ValueError:
        pass
    end = f"i contatti non finiscono il percorso prima del {not_before}" if not_before else ""
    if minutes is None:
        text = "Durata minima del test non nota: la durata di un wait del percorso non è leggibile."
        return text + (f" {end[0].upper()}{end[1:]}." if end else "")
    if minutes > 0:
        return f"Durata minima del test: {_minutes_words(minutes)}" + (f"; {end}." if end else ".")
    if end:
        return f"Durata minima del test: {end}."
    return "Durata minima del test: nessuna, perché i wait del percorso possono terminare subito."


def _action_window(case: Case, action) -> Optional[ActionWindow]:
    """The window coverage.py found for a manual action (``Case.action_windows``)."""
    for window in getattr(case, "action_windows", None) or []:
        if window.activity_key == action.activity_key and window.email_activity_key == action.email_activity_key:
            return window
    return None


# ---------------------------------------------------------------------------
# Test names: a short summary of what makes a case different
# ---------------------------------------------------------------------------

def _short_label(activity: Activity, outcome: Outcome) -> str:
    """An outcome label short enough for a test name. In a two-way split whose condition has a long
    label, the outcomes read "sì" and "no"; any other long label is cut at a word."""
    conditions = [o for o in activity.outcomes if not o.is_remainder]
    remainders = [o for o in activity.outcomes if o.is_remainder]
    if (len(conditions) == 1 and len(remainders) == 1
            and len(conditions[0].label or conditions[0].key) > _LABEL_LIMIT):
        return "sì" if outcome is conditions[0] else "no"
    label = outcome.label or outcome.key
    if len(label) > _LABEL_LIMIT:
        return label[:_LABEL_LIMIT - 3].rsplit(" ", 1)[0] + "..."
    return label


def _name_part(journey: Journey, contact_fields: set[tuple[str, str]], step: Step) -> Optional[tuple[str, int]]:
    """One part of a test name, from a decision or engagement outcome, with its weight: when a name is too
    long, the parts of lowest weight go first. Engagement outcomes and splits on contact fields are what
    make cases of one scenario differ (3); then splits with several values (2); then yes/no switches (1)."""
    activity = journey.activities.get(step.activity_key)
    outcome = activity.outcome(step.outcome_key) if activity is not None and step.outcome_key else None
    if outcome is None:
        return None
    if activity.kind == KIND_ENGAGEMENT:
        if activity.engagement_kind in ("open", "click") and outcome.when is not None:
            email = journey.activities.get(activity.engagement_ref or "")
            name = (email.email_name or email.name) if email is not None else (activity.name or activity.key)
            verb = "aperta" if activity.engagement_kind == "open" else "cliccata"
            return (f"{name} {verb}" if outcome.when else f"{name} non {verb}"), 3
        return f"{activity.name or activity.key}: {outcome.label or outcome.key}", 3
    if activity.kind != KIND_DECISION:
        return None
    text = f"{activity.name or activity.key}: {_short_label(activity, outcome)}"
    if any((c.table.lower(), c.field.lower()) in contact_fields
           for o in activity.outcomes for c in formula_conditions(o.formula)):
        return text, 3
    return text, (2 if len(activity.outcomes) > 2 else 1)


def _name_parts(journey: Journey, contact_fields: set[tuple[str, str]], case: Case) -> list[tuple[str, int]]:
    parts: dict[str, int] = {}
    for step in case.steps:
        part = _name_part(journey, contact_fields, step)
        if part is not None and part[0] not in parts:
            parts[part[0]] = part[1]
    items = list(parts.items())
    if case.random is not None and case.random.branches:
        items.append(("random split " + "/".join(str(b.percentage) for b in case.random.branches), 4))
    return items


def _name_counts(journey: Journey, contact_fields: set[tuple[str, str]], cases: Iterable[Case]) -> dict[str, int]:
    """How many cases use each part of a test name: a part shared by many cases says little."""
    counts: dict[str, int] = {}
    for case in cases:
        for text, _ in _name_parts(journey, contact_fields, case):
            counts[text] = counts.get(text, 0) + 1
    return counts


# ---------------------------------------------------------------------------
# Testbook
# ---------------------------------------------------------------------------

@dataclasses.dataclass
class _TestbookContext:
    result: GenerationResult
    journey: Journey
    settings: Settings
    csv_names: list[str]
    one_active: bool
    activation_ref: Optional[tuple[str, str]]
    first_scenario: Optional[Scenario]
    emails: dict[str, str]           # contact key -> address
    contact_fields: set = dataclasses.field(default_factory=set)    # (table, field) in lower case
    name_counts: dict = dataclasses.field(default_factory=dict)     # test-name part -> number of cases
    entry_kind: str = ENTRY_AUTOMATION                              # EntryEvent.kind
    entry_mode: Optional[str] = None                                # entryMode, when in SINGLE_ENTRY_MODES
    import_names: list = dataclasses.field(default_factory=list)    # CSV files to import (managed tables included)
    managed: list = dataclasses.field(default_factory=list)         # synchronized (managed) tables, noted as such
    real_phone_keys: set = dataclasses.field(default_factory=set)   # contacts that hold the real phone number


def _testbook_context(result: GenerationResult) -> _TestbookContext:
    on, ref, _ = activation_setting(result.package, result.table_model, result.settings)
    scenarios = _ordered_scenarios(result)
    journey = result.package.journey
    contact_fields = {(table.lower(), field.lower()) for table, field in result.coverage.contact_fields}
    managed = _managed_tables(result)
    csv_names = [table_file_name(name) for name in result.data.tables]
    return _TestbookContext(
        result=result,
        journey=journey,
        settings=result.settings,
        csv_names=csv_names,
        import_names=list(csv_names),
        managed=managed,
        real_phone_keys={c.key for c in result.data.contacts if getattr(c, "uses_real_phone", False)},
        one_active=on,
        activation_ref=ref,
        first_scenario=scenarios[0] if scenarios else None,
        emails={contact.key: contact.email for contact in result.data.contacts},
        contact_fields=contact_fields,
        name_counts=_name_counts(journey, contact_fields, result.coverage.cases),
        entry_kind=entry_kind(result.package),
        entry_mode=_entry_mode(result.package),
    )


def build_testbook_rows(result: GenerationResult) -> list[list]:
    """Spec, "Columns of testbook.xlsx": one row per case, scenarios in run order."""
    context = _testbook_context(result)
    rows: list[list] = []
    for scenario in _ordered_scenarios(result):
        for case in scenario.cases:
            rows.append([
                len(rows) + 1,
                OWNER_SYSTEM,
                _test_name(context, case),
                _description(context, scenario, case),
                _preconditions(context, scenario, case),
                _detailed_steps(context, scenario, case),
                _expected_result(context, case),
                STATUS_TO_START,
                _note(context, scenario, case),
            ])
    return rows


def _case_summary(context: _TestbookContext, case: Case) -> str:
    """Spec: "a short summary of the path", for example "Email sent on channel 1, no phone consent".

    A short summary given with the case is kept. A summary that lists every split ("a: b > c: d > ...")
    is rebuilt from the decision and engagement outcomes of the path: when they do not fit, the parts
    that tell least about this case are dropped first (lowest weight, then the part most cases share)."""
    given = (case.summary or "").strip()
    if not given and case.unknown_from is not None:
        return f"Percorso non noto da {_activity_name(context.journey, case.unknown_from)} in poi"
    if not given:
        return f"Percorso fino a {_activity_name(context.journey, case.end_activity_key)}"
    room = _NAME_LIMIT - len(case.id) - len(" - ")
    if " > " not in given and len(given) <= room:
        return given
    parts = _name_parts(context.journey, context.contact_fields, case)
    while len(parts) > 1 and len(", ".join(text for text, _ in parts)) > room:
        drop = min(range(len(parts)), key=lambda i: (parts[i][1], -context.name_counts.get(parts[i][0], 0), i))
        parts.pop(drop)
    text = ", ".join(t for t, _ in parts) if parts else given
    return text if len(text) <= room else text[:room - 3].rstrip() + "..."


def _test_name(context: _TestbookContext, case: Case) -> str:
    """Spec: "Case identifier and a short summary of the path"."""
    return f"{case.id} - {_case_summary(context, case)}"


def _description(context: _TestbookContext, scenario: Scenario, case: Case) -> str:
    """Spec: "Scenario values, contact profile values, and the expected path as activity names in order"."""
    coverage = context.result.coverage
    package = context.result.package
    table, items = _values_parts(package, scenario.values, coverage.scenario_fields)
    if not items:
        scenario_line = f"Scenario {scenario.key}: nessun valore a livello di scenario."
    elif table:
        scenario_line = f"Scenario {scenario.key}, valori in {table}: {'; '.join(items)}."
    else:
        scenario_line = f"Scenario {scenario.key}: {'; '.join(items)}."
    values, rows = _profile(context, case)
    table, items = _values_parts(package, values, coverage.contact_fields)
    rows_text = "; ".join(rows)
    rows_text = rows_text[:1].upper() + rows_text[1:] + "." if rows else ""
    if not items and rows:
        contact_line = f"Profilo del contatto: {'; '.join(rows)}."
    elif not items:
        contact_line = "Profilo del contatto: nessun valore a livello di contatto."
    elif table:
        contact_line = f"Profilo del contatto in {table}: {'; '.join(items)}."
    else:
        contact_line = f"Profilo del contatto: {'; '.join(items)}."
    if items and rows:
        contact_line += f" {rows_text}"
    path = _expected_path_text(context.journey, case.steps, case.unknown_from)
    lines = [scenario_line, contact_line, f"Percorso atteso: {path}."]
    if case.random is not None:
        random_key = case.random.activity_key
        lines.append(f"Al random split {_activity_label(context.journey, random_key)} ogni contatto prende uno dei "
                     "branch.")
        for branch in case.random.branches:
            path = (_expected_path_text(context.journey, branch.steps, branch.unknown_from) if branch.steps
                    else "il contatto esce dal journey")
            lines.append(f"Branch {_branch_label(context.journey, random_key, branch)}: {path}.")
    return "\n".join(lines)


def _absent_tables(context: _TestbookContext, case: Case) -> list[str]:
    """Spec, "Row present or absent": the tables in which the contacts of this case have no row."""
    coverage = context.result.coverage
    return [table for table in getattr(coverage, "row_fields", {}) if coverage.row_absent(case, table)]


def _profile(context: _TestbookContext, case: Case) -> tuple[dict, list[str]]:
    """The contact values to show, and one item per table with the "row present / row absent" variable.

    The field linked to the contact key of such a table says only whether the contact has a row there
    (``GENERATED_KEY`` or None), and every field of an absent row is null: the profile says "riga presente in T"
    or "nessuna riga in T" instead of showing those values."""
    values = dict(case.contact_values)
    rows = []
    for table, linked in getattr(context.result.coverage, "row_fields", {}).items():
        if linked not in values:
            continue
        absent = values[linked] is None
        for ref in [ref for ref in values if ref[0].lower() == table.lower()]:
            if absent or ref == linked:
                values.pop(ref)
        rows.append(f"nessuna riga in {table}" if absent else f"riga presente in {table}")
    return values, rows


def _scenario_key_column(context: _TestbookContext, table_name: str) -> Optional[str]:
    """The field of a table that holds the scenario key."""
    for group in context.result.table_model.key_groups:
        if group.kind != "scenario":
            continue
        for table, field in group.fields:
            if table.lower() == table_name.lower():
                return field
    return None


def _activation_text(context: _TestbookContext, scenario: Scenario) -> str:
    """Spec: "When on, only the first scenario is active in the CSV, and the testbook tells the tester
    which row to switch before each run"; "To move to the next scenario, the tester edits only its
    activation field"."""
    table, field = context.activation_ref
    key_column = _scenario_key_column(context, table)
    row_text = f"riga con {key_column} {scenario.key}" if key_column else f"riga dello scenario {scenario.key}"
    if scenario is context.first_scenario:
        return (f"Verifica che lo scenario {scenario.key} sia quello attivo: in {table}, {field} è True sulla "
                f"{row_text} e False su tutte le altre righe, come generato.")
    return (f"Quando i contatti dello scenario precedente sono usciti dal journey, attiva lo scenario {scenario.key}: "
            f"in {table}, imposta {field} a True sulla {row_text} e a False su tutte le altre righe. Non cambiare "
            "altro.")


def _entry_procedures(result: GenerationResult) -> Optional[str]:
    """The automation that schedules the entry source, when the entry table is also filled by an entry query.
    Only named in the assumption note: the test itself always uses the direct import."""
    entry = result.package.entry_event
    automation = entry.automation_name if entry is not None else None
    return automation if automation and result.table_model.entry_query is not None else None


def _preconditions(context: _TestbookContext, scenario: Scenario, case: Case) -> str:
    """Spec: "Which CSVs must be imported, which scenario must be active, and the keys of the test
    contacts"; "The 50 contacts of a random split case are given as a key range"."""
    entry = context.result.table_model.entry_table
    entry_csv = table_file_name(entry)
    if context.entry_kind == ENTRY_OTHER:
        # Spec: "The team tests a copy of the journey whose entry source is a data extension with the same columns".
        others = [name for name in context.import_names if name != entry_csv]
        first = (f"Importa {entry_csv} nella data extension di ingresso della copia del journey, che ha le stesse "
                 f"colonne di {entry}.")
        if others:
            first = f"Importa {_join_words(others)}, poi " + first[0].lower() + first[1:]
    else:
        first = f"Importa questi file CSV: {', '.join(context.import_names)}."
        if len(context.import_names) > 1 and entry_csv in context.import_names:
            first += f" Importa {entry_csv} per ultimo."
        if context.entry_kind == ENTRY_RUN_ONCE:
            first += " Importa tutti i file prima di attivare il journey."
    lines = [first]
    if context.managed:
        lines.append(_managed_text(context.result, context.managed, short=True))
    if context.one_active:
        table, field = context.activation_ref
        lines.append(f"Scenario attivo: {scenario.key}. In {table}, {field} è True solo sulla sua riga.")
    elif context.entry_kind == ENTRY_RUN_ONCE:
        lines.append(f"Scenario: {scenario.key}. Tutti gli scenari entrano insieme quando il journey viene attivato.")
    else:
        lines.append(f"Scenario: {scenario.key}. Tutti gli scenari sono attivi contemporaneamente.")
    keys = case.contact_keys
    if case.random is not None and len(keys) > 3:
        lines.append(f"Contatti di test ({len(keys)}): da {keys[0]} a {keys[-1]}.")
    else:
        lines.append(f"Contatti di test ({len(keys)}): {', '.join(keys) if keys else 'nessuno'}.")
    absent = _absent_tables(context, case)
    if absent:
        files = _join_words(table_file_name(table) for table in absent)
        lines.append(f"Nessuna riga in {_join_words(absent)} per questi contatti: {files} non contiene le loro chiavi, "
                     "perché il percorso atteso lo richiede.")
    attempt = _attempt_text(context)
    if attempt:
        lines.append(attempt)
    return "\n".join(lines)


def _attempt_text(context: _TestbookContext) -> Optional[str]:
    """Spec, "Entry source and procedure": with Run Once "A new attempt needs a new journey version and a new run
    identifier"; with the entry modes OnceAndDone and SingleEntryAcrossAllVersions "every attempt needs a new run
    identifier"."""
    mode = f" (modalità di ingresso {context.entry_mode})" if context.entry_mode else ""
    if context.entry_kind == ENTRY_RUN_ONCE:
        return f"Ogni nuovo tentativo richiede una nuova versione del journey e un nuovo run identifier{mode}."
    if context.entry_mode:
        return (f"Con la modalità di ingresso {context.entry_mode}, i contatti di un tentativo precedente potrebbero non "
                "rientrare: ogni nuovo tentativo richiede un nuovo run identifier.")
    return None


def _wait_descriptions(context: _TestbookContext, scenario: Scenario, case: Case) -> list[str]:
    """Spec: "Wait by duration, send-time wait: The expected delay is noted in the testbook"."""
    texts = []
    for key in _path_keys(case):
        activity = context.journey.activities.get(key)
        if activity is None:
            continue
        name = _activity_name(context.journey, activity.key)
        if activity.kind == KIND_WAIT_DURATION:
            duration = _duration_text(activity.wait_amount, activity.wait_unit)
            english = _duration_text(activity.wait_amount, activity.wait_unit, english=True)
            texts.append(name if name.lower() in (duration.lower(), english.lower()) else f"{name} ({duration})")
        elif activity.kind == KIND_WAIT_STO:
            if activity.wait_amount:
                texts.append(f"{name}, un send-time wait di al massimo "
                             f"{_duration_text(activity.wait_amount, activity.wait_unit)}")
            else:
                texts.append(f"{name}, un send-time wait la cui attesa non è nota in anticipo")
        elif activity.kind == KIND_WAIT_ATTRIBUTE and activity.wait_field is not None:
            table, field = activity.wait_field
            value = _written_value(context, table, field, scenario, case)
            until = f"fino al {value}, la data in {table}.{field}" if value else f"fino alla data in {table}.{field}"
            texts.append(f"{name}, {until}")
    return texts


def _written_value(context: _TestbookContext, table_name: str, field_name: str,
                   scenario: Scenario, case: Case) -> Value:
    """The value written in the CSV for this scenario's or this case's row of a table."""
    data = next((d for name, d in context.result.data.tables.items() if name.lower() == table_name.lower()), None)
    if data is None or not data.rows:
        return None
    role = next((r for r in context.result.table_model.tables.values() if r.name.lower() == table_name.lower()), None)
    kind = "scenario" if role is not None and role.level == LEVEL_SCENARIO else "contact"
    owner = scenario.key if kind == "scenario" else (case.contact_keys[0] if case.contact_keys else None)
    key_columns = [field for group in context.result.table_model.key_groups if group.kind == kind
                   for table, field in group.fields if table.lower() == table_name.lower()]
    for row in data.rows:
        if any(_row_value(row, column) == owner for column in key_columns):
            return _row_value(row, field_name)
    return None


def _deadline_wait(context: _TestbookContext, scenario: Scenario, case: Case,
                   action) -> tuple[Optional[Activity], Optional[str]]:
    """The last wait by attribute between the email and the engagement split, and its date."""
    for keys in case_paths(case):
        if action.activity_key not in keys:
            continue
        end = keys.index(action.activity_key)
        start = keys.index(action.email_activity_key) if action.email_activity_key in keys[:end] else -1
        for key in reversed(keys[start + 1:end]):
            activity = context.journey.activities.get(key)
            if (activity is not None and activity.kind == KIND_WAIT_ATTRIBUTE and activity.wait_field is not None
                    and wait_precedes_manual_action(key, case)):
                return activity, _written_value(context, activity.wait_field[0], activity.wait_field[1], scenario, case)
    return None, None


def _deadline(context: _TestbookContext, scenario: Scenario, case: Case, action) -> Optional[str]:
    """Spec: "the testbook gives the deadline as a date and time". Spec, "Honest outcomes": "The deadline of a
    manual open or click is computed from the waits between the email and the engagement split": the date of the
    last wait by attribute between them plus the fixed waits after it (``ActionWindow.deadline``). None when the
    window has no wait by attribute: it is then relative to the email (``_relative_window``)."""
    return _deadline_and_wait(context, scenario, case, action)[1]


def _deadline_and_wait(context: _TestbookContext, scenario: Scenario, case: Case,
                       action) -> tuple[Optional[Activity], Optional[str]]:
    """The wait by attribute that ends the window of a manual action, and the deadline."""
    window = _action_window(case, action)
    if window is None or (window.attribute_wait_key is not None and window.deadline is None):
        return _deadline_wait(context, scenario, case, action)
    if window.attribute_wait_key is None:
        return None, None
    return context.journey.activities.get(window.attribute_wait_key), window.deadline


def _relative_window(case: Case, action) -> Optional[ActionWindow]:
    """The window of a manual action that ends a fixed time after the email (no wait by attribute between the
    email and the engagement split, fixed waits of some minutes)."""
    window = _action_window(case, action)
    if window is None or window.attribute_wait_key is not None or not window.minutes:
        return None
    return window


def _manual_steps(context: _TestbookContext, scenario: Scenario, case: Case) -> list[str]:
    """Spec: "any manual action with its deadline"; "Contacts that must open or click get the tester's
    real inbox with plus-addressing"."""
    steps = []
    inbox = context.settings.real_inbox.strip()
    addresses = _join_words(context.emails.get(key, key) for key in case.contact_keys)
    for action in case.manual_actions:
        kind = "open" if action.kind == "open" else "click"
        split = _quoted_label(context.journey, action.activity_key)
        if not inbox:
            steps.append(f"Lo split {split} richiede che i contatti {_ACTION_SUBJUNCTIVE[kind]} '{action.email_name}', "
                         "ma non è stata indicata un'inbox reale, quindi le email non possono essere ricevute. Questo "
                         "test non può raggiungere il risultato atteso finché i dati non vengono generati di nuovo con "
                         "un'inbox reale.")
            continue
        wait, deadline = _deadline_and_wait(context, scenario, case, action)
        window = _action_window(case, action)
        relative = _relative_window(case, action)
        if deadline:
            when = f"Prima del {deadline}"
        elif relative is not None:
            when = f"Entro {_minutes_words(relative.minutes)} dall'arrivo dell'email"
        else:
            when = "Appena arriva l'email"
        text = (f"{when}, {_ACTION_IMPERATIVE[kind]} '{action.email_name}' inviata a ogni contatto di questo test: "
                f"{addresses}.")
        after = window.minutes_after_attribute_wait if window is not None else 0
        if deadline and wait is not None and not after:
            text += (f" A quell'ora termina il wait {_quoted_label(context.journey, wait.key)} e i contatti "
                     f"proseguono verso lo split {split}.")
        elif deadline and wait is not None:
            text += (f" A quell'ora i contatti hanno superato il wait {_quoted_label(context.journey, wait.key)} e i "
                     f"wait successivi ({_minutes_words(after)}) e raggiungono lo split {split}.")
        elif relative is not None:
            text += f" Dopo questo tempo i contatti raggiungono lo split {split}."
        steps.append(text + f" Le email arrivano nell'inbox {inbox}. Non aprire e non cliccare le email degli altri "
                     "test.")
    return steps


def _detailed_steps(context: _TestbookContext, scenario: Scenario, case: Case) -> str:
    """Spec: "Numbered steps: activate the scenario, start the entry, wait (with expected delays), any
    manual action with its deadline, collect evidence"."""
    steps = []
    entry = context.result.table_model.entry_table
    if context.one_active:
        steps.append(_activation_text(context, scenario))
    elif context.entry_kind == ENTRY_RUN_ONCE:
        steps.append(f"Verifica che le righe dello scenario {scenario.key} siano state importate e che {entry} "
                     "contenga solo righe di test.")
    else:
        steps.append(f"Verifica che le righe dello scenario {scenario.key} siano state importate.")
    steps.append(_entry_step(context))
    timeline = _timeline_step(context, case)
    if timeline:
        steps.append(timeline)
    waits = _wait_descriptions(context, scenario, case)
    unknown = _unknown_from(case)
    if waits:
        text = ("Attendi che i contatti superino i wait del loro percorso: " + "; ".join(waits) + ". "
                + _duration_sentence(case, context.result.generation_time, context.settings.date_format))
        if unknown is not None:
            # Spec, "Honest outcomes": past an unsupported element the path, and so its waits, are unknown.
            text += (f" La durata vale solo per la parte nota del percorso, fino a "
                     f"{_activity_label(context.journey, unknown)}: dopo, il percorso non è noto e può avere altri wait.")
        steps.append(text)
    elif unknown is not None:
        steps.append(f"La parte nota del percorso, fino a {_activity_label(context.journey, unknown)}, non ha wait: "
                     "dopo, il percorso non è noto e può avere dei wait.")
    else:
        steps.append("Il percorso non ha wait.")
    steps.extend(_manual_steps(context, scenario, case))
    steps.extend(_sms_steps(context, case))
    steps.append("Raccogli le evidenze per ogni contatto: le email ricevute, i valori scritti dal journey e l'ultima "
                 "attività raggiunta.")
    return "\n".join(f"{number}. {text}" for number, text in enumerate(steps, 1))


def _timeline_text(context: _TestbookContext, point, exit_key: Optional[str]) -> str:
    """One step of the timeline in words; the wait at whose end the exit criteria take the contact out says so."""
    journey = context.journey
    activity = journey.activities.get(point.activity_key)
    label = _quoted_label(journey, point.activity_key)
    if activity is None:
        return label
    kind = activity.kind
    if kind in (KIND_WAIT_DURATION, KIND_WAIT_ATTRIBUTE, KIND_WAIT_STO):
        if exit_key == point.activity_key and point.outcome_key is None:
            return f"il contatto esce dal journey alla fine del wait {label} per i criteri di uscita{_exit_phrase(journey)}"
        return f"fine del {'send-time wait' if kind == KIND_WAIT_STO else 'wait'} {label}"
    if kind == KIND_EMAIL:
        return f"invio dell'email {label}"
    if kind == KIND_SMS:
        return f"invio dell'SMS {label}"
    if kind == KIND_UPDATE:
        return f"Update contact {label}"
    if point.outcome_key is not None and len(activity.outcomes) > 1:
        return f"{label}, branch '{_outcome_label(journey, point.activity_key, point.outcome_key)}'"
    return label


def _timeline_lines(context: _TestbookContext, points, exit_key: Optional[str]) -> list[str]:
    """Spec, step 3, "The clock": "The testbook writes it as days from the entry time with the absolute date and
    time." One line per moment: the steps that happen at the same moment share a line."""
    start = entry_time(context.result)
    lines: list[str] = []
    group: list[str] = []
    current = None
    for point in points:
        if group and point.at != current:
            lines.append(f"Giorno {(current - start).days} ({_date(current, context.settings)}): {'; '.join(group)}")
            group = []
        current = point.at
        group.append(_timeline_text(context, point, exit_key))
    if group:
        lines.append(f"Giorno {(current - start).days} ({_date(current, context.settings)}): {'; '.join(group)}")
    return lines


def _timeline_step(context: _TestbookContext, case: Case) -> Optional[str]:
    """The timeline of the case as one numbered step; nothing for a case without one."""
    timeline = list(getattr(case, "timeline", None) or [])
    if not timeline:
        return None
    lines = ["   " + line for line in _timeline_lines(context, timeline, case.exit_by_criteria)]
    if case.random is not None:
        for branch in case.random.branches:
            branch_lines = _timeline_lines(context, getattr(branch, "timeline", ()) or (), branch.exit_by_criteria)
            if branch_lines:
                lines.append(f"   Branch {_branch_label(context.journey, case.random.activity_key, branch)} del random split:")
                lines.extend("      " + line for line in branch_lines)
    start = _date(entry_time(context.result), context.settings)
    return f"Controlla che il percorso rispetti questi tempi, contati dall'ora di ingresso ({start}):\n" + "\n".join(lines)


def _sms_steps(context: _TestbookContext, case: Case) -> list[str]:
    """Spec, step 3, "SMS sends": the step names the SMS and the contact that receives it (the one with the real
    number); without a real number the SMS is sent to the fake one and cannot be checked."""
    settings = context.settings
    receiving = [key for key in case.contact_keys if key in context.real_phone_keys]
    mobile = getattr(context.journey, "default_mobile_field", None)
    where = f" nella colonna {mobile} di {context.result.table_model.entry_table}" if mobile else ""
    steps = []
    for key in _unique(getattr(case, "sms_sent", None) or []):
        label = _quoted_label(context.journey, key)
        if receiving:
            steps.append(f"SMS {label} inviato al numero del contatto {receiving[0]}, che ha il numero reale "
                         f"{settings.real_phone.strip()}{where}. Controlla che l'SMS arrivi sul telefono.")
        else:
            steps.append(f"SMS {label} inviato ai contatti del test, che hanno il numero fittizio "
                         f"{settings.fake_phone}: senza un numero di telefono reale l'SMS non può essere verificato. "
                         "Controlla solo l'ultima attività raggiunta.")
    return steps


def _sms_result(context: _TestbookContext, case: Case) -> Optional[str]:
    """The SMS received, in the expected result."""
    sent = _unique(getattr(case, "sms_sent", None) or [])
    if not sent:
        return None
    names = _join_words(_quoted_label(context.journey, key) for key in sent)
    receiving = [key for key in case.contact_keys if key in context.real_phone_keys]
    if receiving:
        return (f"- SMS ricevuti: {names} sul telefono del contatto {receiving[0]}, che ha il numero reale "
                f"{context.settings.real_phone.strip()}; gli altri contatti hanno il numero fittizio "
                f"{context.settings.fake_phone}.")
    return (f"- SMS inviati: {names}; non verificabile senza numero reale, perché i contatti hanno il numero fittizio "
            f"{context.settings.fake_phone}.")


def _exit_result(context: _TestbookContext, wait_key: str) -> str:
    """The exit by criteria as the last event of the expected result."""
    return (f"i contatti escono dal journey alla fine del wait {_quoted_label(context.journey, wait_key)}, perché "
            f"soddisfano i criteri di uscita{_exit_phrase(context.journey)}.")


def _unknown_from(case: Case) -> Optional[str]:
    """The activity from which the path of the case, or of one of its random branches, is unknown."""
    if case.unknown_from is not None:
        return case.unknown_from
    branches = case.random.branches if case.random is not None else []
    return next((b.unknown_from for b in branches if getattr(b, "unknown_from", None) is not None), None)


def _entry_step(context: _TestbookContext) -> str:
    """How the contacts enter, by entry source (spec, "Entry source and procedure")."""
    name = context.journey.name
    entry = context.result.table_model.entry_table
    entry_csv = table_file_name(entry)
    if context.entry_kind == ENTRY_RUN_ONCE:
        return (f"Attiva il journey {name}: entrano solo le righe presenti in {entry} in quel momento, e le righe "
                "importate dopo non entrano. Tutti i test partono insieme.")
    if context.entry_kind == ENTRY_SCHEDULED:
        return (f"Attendi la valutazione successiva della data extension {entry}: i contatti dovrebbero entrare nel "
                f"journey {name} in quel momento (procedura da verificare). Tutti i test di questo scenario partono "
                "insieme.")
    if context.entry_kind == ENTRY_OTHER:
        return (f"Fai entrare i contatti nella copia del journey {name}, con entry source la data extension di "
                f"ingresso, dalle righe importate di {entry_csv}. Tutti i test di questo scenario partono insieme.")
    return (f"Avvia l'ingresso nel journey {name} dalle righe importate di {entry_csv}. Tutti i test di questo "
            "scenario partono insieme.")


def _expected_result(context: _TestbookContext, case: Case) -> str:
    """Spec: "For each contact: the emails or messages received, the values written by the journey, and
    the last activity reached". The contacts of a case share one expected result."""
    journey = context.journey
    keys = case.contact_keys
    now_text = _date(context.result.generation_time, context.settings)
    base_keys = [step.activity_key for step in case.steps]
    if case.random is None and case.unknown_from is not None:
        # Spec, "Honest outcomes": the path is unknown from that activity on, never "the last activity reached".
        before = f" prima di {_activity_label(journey, case.unknown_from)}"
        sms = _sms_result(context, case)
        return "\n".join([
            f"Per ciascuno di questi contatti: {_join_words(keys) if keys else 'nessuno'}.",
            f"- Email ricevute{before}: {_emails_text(journey, case.emails_sent)}.",
            f"- Valori scritti dal journey{before}: "
            f"{_writes_text(case.writes, _current_time_fields(journey, base_keys), now_text)}.",
        ] + ([sms] if sms else []) + [
            f"- {_unknown_text(journey, case.unknown_from, case.unknown_reason)}. Annota l'ultima attività raggiunta "
            "da ogni contatto.",
        ])
    if case.random is None:
        last = f"- Ultima attività raggiunta: {_activity_label(journey, case.end_activity_key)}."
        end = _end_sentence(journey, case.steps[-1] if case.steps else None)
        sms = _sms_result(context, case)
        lines = [
            f"Per ciascuno di questi contatti: {_join_words(keys) if keys else 'nessuno'}.",
            f"- Email ricevute: {_emails_text(journey, case.emails_sent)}.",
            f"- Valori scritti dal journey: "
            f"{_writes_text(case.writes, _current_time_fields(journey, base_keys), now_text)}.",
        ] + ([sms] if sms else []) + [f"{last} {end}" if end else last]
        if case.exit_by_criteria:
            # Spec, step 3: the expected path ends with "leaves the journey at the end of <wait>, because the
            # exit criteria are met".
            lines.append("- Uscita: " + _exit_result(context, case.exit_by_criteria))
        return "\n".join(lines)
    random_key = case.random.activity_key
    lines = [
        f"Per ciascuno dei {len(keys)} contatti {_key_range(keys)}, fino al random split "
        f"{_activity_label(journey, random_key)}:",
        f"- Email ricevute: {_emails_text(journey, case.emails_sent)}.",
        f"- Valori scritti dal journey: "
        f"{_writes_text(case.writes, _current_time_fields(journey, base_keys), now_text)}.",
    ]
    sms = _sms_result(context, case)
    if sms:
        lines.append(sms)
    lines.append("Poi ogni contatto prende uno dei branch del random split:")
    for branch in case.random.branches:
        branch_keys = base_keys + [step.activity_key for step in branch.steps]
        if branch.unknown_from is not None:
            last = _unknown_text(journey, branch.unknown_from, branch.unknown_reason)
            last = last[0].lower() + last[1:]
        else:
            last = f"ultima attività raggiunta: {_activity_label(journey, branch.end_activity_key)}"
        if getattr(branch, "exit_by_criteria", None):
            last += "; " + _exit_result(context, branch.exit_by_criteria).rstrip(".")
        lines.append(f"- {_branch_label(journey, random_key, branch)}: email ricevute: "
                     f"{_emails_text(journey, branch.emails_sent)}; "
                     f"valori scritti: {_writes_text(branch.writes, _current_time_fields(journey, branch_keys), now_text)}; "
                     f"{last}.")
    smallest = _smallest_branch(case)
    if smallest is not None:
        lines.append(f"Intervallo accettato: da {case.random.expected_low} a {case.random.expected_high} contatti sul "
                     f"branch {_branch_label(journey, random_key, smallest)}.")
    return "\n".join(lines)


def _note(context: _TestbookContext, scenario: Scenario, case: Case) -> str:
    """Spec: "Manual engagement required, or the accepted range for a random split case"."""
    notes = []
    for action in case.manual_actions:
        kind = "open" if action.kind == "open" else "click"
        verb = _ACTION_TO_DO[kind]
        if not context.settings.real_inbox.strip():
            notes.append(f"Azione manuale richiesta: {verb} '{action.email_name}'. Non è stata indicata un'inbox "
                         "reale, quindi questo test non può raggiungere il risultato atteso.")
            continue
        wait, deadline = _deadline_and_wait(context, scenario, case, action)
        window = _action_window(case, action)
        relative = _relative_window(case, action)
        text = f"Azione manuale richiesta: {verb} '{action.email_name}'"
        if deadline:
            notes.append(text + f" prima del {deadline}.")
        elif relative is not None:
            notes.append(text + f" entro {_minutes_words(relative.minutes)} dall'arrivo.")
        else:
            notes.append(text + " appena arriva.")
        if deadline:
            notes.append("Se l'email arriva dopo quell'ora, genera di nuovo i dati.")
        if window is not None and window.short:
            # Spec, "Honest outcomes": "with a warning when the window is under one hour".
            arrival = getattr(window, "email_not_before", None)
            if deadline and arrival:
                # The waits before the email delay it: the window starts when it arrives.
                notes.append(f"L'email arriva non prima del {arrival}, quindi la scadenza è a meno di un'ora dal suo "
                             "arrivo: genera di nuovo i dati con un offset della data di wait prima di un engagement "
                             "split più ampio.")
            elif deadline:
                notes.append("La scadenza è a meno di un'ora dall'ora di generazione: se non c'è tempo, genera di nuovo "
                             "i dati con un offset della data di wait prima di un engagement split più ampio.")
            else:
                notes.append(f"La finestra è sotto l'ora: {_ACTION_IMPERATIVE[kind]} appena arriva.")
        zone = wait.wait_timezone if deadline and wait is not None else None
        if zone:
            notes.append(f"Il wait {_quoted_label(context.journey, wait.key)} è impostato sul fuso orario {zone}. "
                         "Come Salesforce applica questo fuso orario alla data è un'assunzione da verificare, quindi "
                         "la scadenza potrebbe spostarsi di qualche ora.")
    if case.random is not None:
        smallest = _smallest_branch(case)
        label = _branch_label(context.journey, case.random.activity_key, smallest) if smallest else "più piccolo"
        contacts = len(case.contact_keys)
        notes.append(f"Random split: non è possibile indirizzare i contatti su un branch preciso. Sul branch {label} "
                     f"sono attesi da {case.random.expected_low} a {case.random.expected_high} dei {contacts} "
                     "contatti.")
        chance = _range_chance(contacts, smallest.percentage if smallest else None,
                               case.random.expected_low, case.random.expected_high)
        if chance is not None:
            notes.append(f"Con una probabilità del {smallest.percentage}% per ogni contatto, questo intervallo si "
                         f"verifica in circa il {chance}% dei run.")
    notes.extend(_time_notes(context, case))
    not_modelled = _not_modelled_text(context.result.package)
    if not_modelled:
        notes.append(not_modelled)
    return " ".join(notes)


def _time_notes(context: _TestbookContext, case: Case) -> list[str]:
    """Spec, step 3: the deadline of this case ("Validity of the data") and, when a branch of its path rests on a
    bound, the day-level comparison it is decided on (the note coverage.py wrote for this case)."""
    notes = []
    days = getattr(case, "validity_days", None)
    if days is not None:
        result = context.result
        start = entry_time(result)
        notes.append(f"Attiva entro {_date(start + timedelta(days=days), context.settings)}: i valori di questo test "
                     f"restano validi {_plural(days, 'giorno', 'giorni')} dall'ora di ingresso "
                     f"({_date(start, context.settings)}), poi un branch o un wait potrebbe leggerli in modo diverso.")
    for warning in context.result.coverage.warnings:
        if _DAY_LEVEL_WORDS not in warning:
            continue
        named = re.split(r"[\s,]+", warning.split(", in cui")[0])
        if case.id in named:
            notes.append(_sentence(warning))
    return notes


def _excel_text(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _utc_text(moment: datetime) -> str:
    if moment.tzinfo is not None:
        moment = moment.astimezone(timezone.utc)
    return moment.strftime("%Y-%m-%dT%H:%M:%SZ")


def _stable_xlsx(content: bytes, moment: datetime) -> bytes:
    """openpyxl stamps the save time into docProps/core.xml and the zip entries: replace both, so the
    same input always gives the same bytes."""
    stamp = _utc_text(moment)
    source = zipfile.ZipFile(io.BytesIO(content))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                text = _CORE_TIMESTAMP.sub(lambda m: m.group(1) + stamp + m.group(2), data.decode("utf-8"))
                data = text.encode("utf-8")
            archive.writestr(_zip_info(info.filename), data)
    return buffer.getvalue()


def write_testbook(result: GenerationResult) -> bytes:
    """Spec, "Outputs": "testbook.xlsx: One test per case with the expected result, in the team's
    testbook layout"; "The testbook has one row per case, not per contact"."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = TESTBOOK_SHEET
    sheet.append(TESTBOOK_COLUMNS)
    for row in build_testbook_rows(result):
        sheet.append([_excel_text(value) for value in row])
    header_fill = PatternFill("solid", fgColor="1F3864")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for position, width in enumerate(_TESTBOOK_WIDTHS, 1):
        sheet.column_dimensions[sheet.cell(1, position).column_letter].width = width
    sheet.freeze_panes = "A2"
    workbook.properties.creator = "CJ Test Helper"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return _stable_xlsx(buffer.getvalue(), result.generation_time)


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------

def _md_cell(value) -> str:
    text = "" if value is None else str(value)
    return text.replace("|", "\\|").replace("\r", " ").replace("\n", " ")


def _md_table(header: list[str], rows: list[list]) -> str:
    lines = ["| " + " | ".join(_md_cell(h) for h in header) + " |",
             "| " + " | ".join("---" for _ in header) + " |"]
    lines += ["| " + " | ".join(_md_cell(value) for value in row) + " |" for row in rows]
    return "\n".join(lines)


def _bullets(items: Iterable[str]) -> str:
    return "\n".join(f"- {item}" for item in items)


def _numbered(items: Iterable[str]) -> str:
    return "\n".join(f"{number}. {item}" for number, item in enumerate(items, 1))


def _unique(items: Iterable[str]) -> list[str]:
    seen: list[str] = []
    for item in items:
        if item not in seen:
            seen.append(item)
    return seen


def _split_suspicious(warnings: Iterable[str]) -> tuple[list[str], list[str]]:
    """Spec, "Outputs": the report "lists suspicious patterns found while parsing, such as the same value
    written in two letter cases". Other modules mark them with the prefix ``pattern sospetto:``."""
    patterns, others = [], []
    for warning in warnings:
        if warning.lower().startswith(SUSPICIOUS_PREFIX):
            pattern = warning[len(SUSPICIOUS_PREFIX):].strip()
            patterns.append(pattern[:1].upper() + pattern[1:])
        else:
            others.append(warning)
    return patterns, others


# Texts written by other modules that the report rewords. A text that no longer matches is shown as it is.
# Suspicious patterns arrive with their first letter in upper case (``_split_suspicious``), hence ``[Ll]``.
_CASE_PATTERN = re.compile(r"^(?P<table>[^.]+)\.(?P<field>.+?) è confrontato con (?P<values>.+?)"
                           r"(?:; (?P<tail>.+?))?\.?$")
_SAME_NAME_PATTERN = re.compile(r"^(?P<count>\d+) split con condizioni diverse hanno lo stesso nome '(?P<name>.*)' "
                                r"\((?P<keys>.*)\)\.?$")
_CONSTANT_KEY_PATTERN = re.compile(r"^[Ll]a costante (?P<value>.+?) è scritta nel campo di chiave primaria "
                                   r"(?P<table>[^.]+)\.(?P<field>.+?) da (?P<count>\d+) attività Update contact, "
                                   r"quindi tutti i contatti hanno lo stesso valore di chiave in quel campo "
                                   r"\((?P<keys>.*)\)\.?$")
_ONE_TO_ONE_PATTERN = re.compile(r"^[Ll]a relazione uno a uno '(?P<name>.*)' collega (?P<table>.+?) tramite "
                                 r"(?P<fields>.+?), ma la chiave primaria della tabella comprende anche (?P<extra>.+?): "
                                 r"a un contatto possono corrispondere più righe\.?$")
_CASE_TAILS = {
    "sono valori diversi perché i confronti di testo distinguono maiuscole e minuscole (impostazione)":
        "Con l'impostazione attuale i confronti di testo distinguono maiuscole e minuscole, quindi sono valori diversi.",
    "sono trattati come un solo valore perché i confronti di testo non distinguono maiuscole e minuscole":
        "Con l'impostazione attuale i confronti di testo non distinguono maiuscole e minuscole, quindi contano come "
        "un solo valore.",
}
_IGNORED_COLUMNS = re.compile(r"^La entry query seleziona colonne che (?P<table>.+?) non ha\. Queste colonne "
                              r"vengono ignorate: (?P<columns>.+?)\.?$")
_QUOTED = re.compile(r"'((?:[^'\\]|\\.)*)'|\"((?:[^\"\\]|\\.)*)\"|\b(null)\b")
_UNREACHABLE_PREFIX = "non raggiungibile: "
_REASON_NEVER_MATCHES = re.compile(r"^non raggiungibile: lo split è raggiunto, ma su quei percorsi la condizione di "
                                   r"'(?P<label>.*)' non è mai vera \((?P<detail>.*)\)$")
_REASON_ALWAYS_ONE_OF = re.compile(r"^(?P<field>.+?) vale sempre uno tra (?P<values>.+)$")
_REASON_NEVER_TRUE = "la condizione non è mai vera"
_REASON_FOLLOWS = re.compile(r"^non raggiungibile: nessuna combinazione di valori dei campi raggiunge questa attività "
                             r"\(segue (?P<keys>.*)\)$")
_REASON_BEHIND_UNSUPPORTED = re.compile(r"^a valle di un elemento non supportato: (?P<detail>.+)$")
_REASON_AFTER_UNSUPPORTED = re.compile(r"^questo branch viene dopo il branch non supportato '(?P<label>.*)' dello "
                                       r"stesso split, quindi non viene mai valutato$")
_REASON_UNSUPPORTED = re.compile(r"^(?P<what>attività non supportata|branch non supportato): (?P<detail>.+)$")
_REASON_BEHIND_BRANCH = re.compile(r"^(?:non raggiungibile|a valle di un elemento non supportato): si trova a valle "
                                   r"del branch .*?, che non è coperto per questo motivo: (?P<reason>.+)$")
_REASON_BLOCKED = re.compile(r"^richiede un open o un click che non si può fare: (?P<detail>.+); i casi esistono$")
_REASON_FIXED = {
    "non raggiungibile: lo split è raggiunto, ma la condizione di un branch precedente è sempre vera":
        ("unreachable", "la condizione di un branch precedente dello split è sempre vera"),
    "non raggiungibile: lo split è raggiunto, ma nessun contatto prende questo branch":
        ("unreachable", "nessun contatto prende questo branch, anche se lo split viene raggiunto"),
    "non raggiungibile: nessun percorso simulato attraversa questo branch":
        ("unreachable", "nessun percorso simulato lo attraversa"),
    "non raggiungibile: nessun percorso dall'inizio del journey raggiunge questa attività":
        ("unreachable", "nessun percorso dall'inizio del journey raggiunge lo split"),
    "richiede un open o un click, ma non è configurata un'inbox reale; i casi esistono, con indirizzi non recapitabili":
        ("not covered", "richiede un open o un click, ma non è stata indicata un'inbox reale. I suoi casi esistono, "
                        "con indirizzi non recapitabili"),
}


def _quoted_values(text: str) -> list[str]:
    """The values of a list such as ``'DM', 'Email', null``, as they should be shown."""
    shown = []
    for single, double, null in _QUOTED.findall(text):
        shown.append("vuoto" if null else (f"'{single}'" if single or not double else f'"{double}"'))
    return shown


def _or_values(text: str) -> str:
    """``'DM', 'Email', 'OB'`` as ``'DM', 'Email' o 'OB'``."""
    values = _quoted_values(text)
    if not values:
        return "uno tra " + text
    return values[0] if len(values) == 1 else ", ".join(values[:-1]) + " o " + values[-1]


def _quoted_labels(journey: Journey, text: str) -> list[str]:
    """Activity keys written as ``'A', 'B'``, by name and key."""
    keys = [single or double for single, double, _ in _QUOTED.findall(text) if single or double]
    return [_activity_label(journey, key) for key in keys] or [text]


def _drop_table(detail: str, condition: str) -> str:
    """``T.F vale sempre ...`` as ``F vale sempre ...`` when the quoted condition already names ``T.F``."""
    field = detail.split(" vale sempre ", 1)[0]
    if "." in field and field in condition:
        return field.split(".", 1)[1] + detail[len(field):]
    return detail


def _natural_key(text: str) -> list:
    """Sort key that puts MULTICRITERIADECISIONV2-42 before MULTICRITERIADECISIONV2-116."""
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", text)]


def _condition_text(journey: Journey, activity_key: str, outcome_key: str) -> Optional[str]:
    activity = journey.activities.get(activity_key)
    outcome = activity.outcome(outcome_key) if activity is not None else None
    if outcome is None or outcome.formula is None:
        return None
    text = formula_text(outcome.formula)
    return f"`{text}`" if len(text) <= _FORMULA_LIMIT else None


def _explain_reason(reason: str, journey: Journey, activity_key: str,
                    outcome_key: str) -> tuple[Optional[str], str]:
    """A not-covered reason from the coverage module as (state, clause), where state is "unreachable",
    "not covered", or None when the reason is not one the report knows (it is then shown as it is)."""
    reason = (reason or "").strip()
    if reason in _REASON_FIXED:
        return _REASON_FIXED[reason]
    match = _REASON_NEVER_MATCHES.match(reason)
    if match:
        details = []
        for part in match.group("detail").split("; "):
            seen = _REASON_ALWAYS_ONE_OF.match(part)
            if seen:
                details.append(f"{seen.group('field')} vale sempre {_or_values(seen.group('values'))}")
            elif part != _REASON_NEVER_TRUE:
                details.append(part)
        condition = _condition_text(journey, activity_key, outcome_key)
        if condition:
            details = [_drop_table(detail, condition) for detail in details]
        head = (f"la sua condizione {condition} non è mai soddisfatta" if condition
                else f"'{match.group('label')}' non è mai soddisfatto")
        if details:
            return "unreachable", f"{head}: dove lo split viene raggiunto, {_join_words(details)}"
        return "unreachable", f"{head} dove lo split viene raggiunto"
    match = _REASON_FOLLOWS.match(reason)
    if match:
        parents = _join_words(_quoted_labels(journey, match.group("keys")))
        return "unreachable", f"nessuna combinazione di valori dei campi raggiunge lo split, che segue {parents}"
    match = _REASON_BEHIND_UNSUPPORTED.match(reason)
    if match:
        return "not covered", match.group("detail")
    match = _REASON_UNSUPPORTED.match(reason)
    if match:
        what = ("lo split stesso non è supportato" if match.group("what").startswith("attività")
                else "il branch non è supportato")
        return "not covered", f"{what}: {match.group('detail').rstrip('.')}"
    match = _REASON_BLOCKED.match(reason)
    if match:
        return "not covered", (f"richiede un open o un click che non si può fare: {match.group('detail')}. I suoi "
                               "casi esistono")
    if reason.startswith(_UNREACHABLE_PREFIX):
        return "unreachable", reason[len(_UNREACHABLE_PREFIX):]
    return None, reason


def _reason_sentence(reason: str, journey: Journey, activity_key: str, outcome_key: str) -> str:
    state, clause = _explain_reason(reason, journey, activity_key, outcome_key)
    if state == "unreachable":
        return _sentence(f"Non raggiungibile perché {clause.rstrip('.')}")
    if state == "not covered":
        return _sentence(f"Non coperto perché {clause.rstrip('.')}")
    return _sentence(reason) if reason else "Nessun motivo indicato."


def _section_title(result: GenerationResult) -> str:
    settings = result.settings
    journey = result.package.journey.name
    intro = "Questo report indica cosa coprono i dati di test, come eseguire i test e cosa i dati non possono coprire."
    if TESTBOOK_FILE in result.files:
        intro += f" I passi e il risultato atteso di ogni test sono in {TESTBOOK_FILE}."
    parts = [
        f"# Report dei dati di test: {journey}",
        f"Journey {journey}, dal package {result.package.name}. Generato il "
        f"{_date(result.generation_time, settings)} con il run identifier {settings.run_id}.",
        intro,
    ]
    deadline = activation_deadline_text(result)
    if deadline:
        # Spec, step 3, "Honest checks": "The report opens with the deadline for activation when the data is
        # perishable, next to the warning on unsupported elements."
        parts.append(f"> **Attenzione:** {deadline}")
    warning = unsupported_warning(result)
    if warning:
        # Spec, "Honest outcomes": "report.md and the page open with a warning when unsupported elements leave
        # branches not covered."
        parts.append(f"> **Attenzione:** {warning} Le sezioni Branch non coperti ed Elementi non supportati spiegano "
                     "il motivo.")
    return "\n\n".join(parts)


def _section_summary(result: GenerationResult) -> str:
    coverage, data = result.coverage, result.data
    cases = coverage.cases
    statuses = [b.status for b in coverage.branches]
    covered = statuses.count(STATUS_COVERED)
    manual = statuses.count(STATUS_MANUAL)
    not_covered = statuses.count(STATUS_NOT_COVERED)
    statistical = sum(1 for b in coverage.branches if b.statistical and b.status != STATUS_NOT_COVERED)
    rows = [
        ["Scenari", len(coverage.scenarios)],
        ["Scenari statistici", sum(1 for s in coverage.scenarios if s.is_statistical)],
        ["Casi", len(cases)],
        ["Casi con un'azione manuale", sum(1 for c in cases if c.manual_actions)],
        ["Casi statistici, per un random split", sum(1 for c in cases if c.random is not None)],
        ["Contatti", len(data.contacts)],
        ["Contatti con l'inbox reale", sum(1 for c in data.contacts if c.uses_real_inbox)],
        ["Branch", len(coverage.branches)],
        ["Branch coperti", covered],
        ["Branch coperti con un'azione manuale", manual],
        ["Branch non coperti", not_covered],
        ["Branch coperti solo tramite un random split", statistical],
    ]
    targets = list(getattr(coverage, "exit_targets", None) or [])
    exits_covered = sum(1 for t in targets if t.status != STATUS_NOT_COVERED)
    if targets:
        # Spec, step 3: "the summary counts it apart from the branches of the splits".
        rows += [["Target dei criteri di uscita", len(targets)], ["Target dei criteri di uscita coperti", exits_covered]]
    text = (f"I dati di test comprendono {_plural(len(coverage.scenarios), 'scenario', 'scenari')}, "
            f"{_plural(len(cases), 'caso', 'casi')} e {_plural(len(data.contacts), 'contatto', 'contatti')}. "
            f"Su {len(coverage.branches)} branch, {_covered(covered)}, "
            f"{_covered(manual)} con un'azione manuale e {_covered(not_covered, 'non coperto')}.")
    if targets:
        text += (f" Contati a parte dai branch degli split, i criteri di uscita (exit criteria) hanno "
                 f"{_plural(len(targets), 'target', 'target')}: {exits_covered} "
                 f"{'coperto' if exits_covered == 1 else 'coperti'}.")
    tables = list(result.table_model.scenario_tables)
    if len(tables) == 1:
        terms = f"Uno scenario è una riga di {tables[0]}, condivisa da tutti i suoi contatti."
    elif tables:
        terms = (f"Uno scenario è una riga in ciascuna delle tabelle {_join_words(tables)}, condivisa da tutti i "
                 "suoi contatti.")
    else:
        terms = "Uno scenario raggruppa i casi che vengono eseguiti insieme."
    terms += (" Un caso raggruppa contatti con gli stessi valori di contatto e le stesse azioni manuali, che quindi "
              "seguono lo stesso percorso fino all'eventuale random split.")
    if TESTBOOK_FILE in result.files:
        terms += f" Ogni caso è un test di {TESTBOOK_FILE}."
    return "\n\n".join(["## Riepilogo", text, terms, _md_table(["Voce", "Numero"], rows)])


def _section_checks(result: GenerationResult) -> str:
    if not result.checks:
        return "## Trust check\n\nSu questi file non è stato eseguito alcun trust check."
    failed = [check for check in result.checks if not check.passed]
    total = len(result.checks)
    if failed:
        passed = total - len(failed)
        text = (f"{passed} trust check su {total} {'è superato' if passed == 1 else 'sono superati'}, "
                f"{len(failed)} {'non è superato' if len(failed) == 1 else 'non sono superati'}: i dettagli sono "
                "qui sotto.")
    else:
        text = "Il trust check è superato." if total == 1 else f"Tutti i {total} trust check sono superati."
    meaning = ("Dopo aver scritto i file, il generatore li rilegge e li controlla. Un check non superato significa che "
               "i file potrebbero non fare quello che dicono questo report e il testbook. Non importarli finché la "
               "causa non è chiarita.")
    rows = [[check.name, "superato" if check.passed else "NON SUPERATO", len(check.details)]
            for check in result.checks]
    parts = ["## Trust check", text, meaning, _md_table(["Trust check", "Esito", "Dettagli"], rows)]
    for check in result.checks:
        if not check.details:
            continue
        shown = check.details[:_REPORT_DETAIL_LIMIT]
        lines = [f"### {check.name}", _bullets(_sentence(detail) for detail in shown)]
        if len(check.details) > len(shown):
            lines.append(f"Ce ne sono altri {len(check.details) - len(shown)}.")
        parts.append("\n\n".join(lines))
    return "\n\n".join(parts)


def _section_tables(result: GenerationResult) -> str:
    rows = []
    managed = _managed_tables(result)
    for role in result.table_model.tables.values():
        data = result.data.tables.get(role.name)
        csv_name = table_file_name(role.name) if role.selected else "no"
        rows.append([role.name, role_text(role.roles),
                     _LEVEL_WORDS.get(role.level, role.level) if role.level else "-",
                     csv_name, len(data.rows) if data is not None else "-", role.reason])
    selected = [table_file_name(name) for name in result.data.tables]
    text = (f"{_plural(len(selected), 'tabella riceve', 'tabelle ricevono')} un CSV: "
            f"{_join_words(selected) if selected else 'nessuna'}.")
    meaning = ("Una tabella riceve un CSV quando il journey o la entry query la leggono. Una tabella a livello di "
               "contatto ha una riga per contatto. Una tabella a livello di scenario ha una riga per scenario, "
               "condivisa dai suoi contatti.")
    if any(ROLE_LINK in role.roles for role in result.table_model.tables.values()):
        meaning += (" Una tabella di collegamento sta sul percorso tra il contatto e una tabella che uno split o un "
                    "wait legge.")
    parts = ["## Tabelle", text, meaning, _md_table(["Tabella", "Ruolo", "Livello", "CSV", "Righe", "Motivo"], rows)]
    if managed:
        parts.append(_managed_text(result, managed))
    return "\n\n".join(parts)


def _activation_is_entry_filter(result: GenerationResult, ref: FieldRef) -> bool:
    """True when the activation field is a filter of the entry query and no split or wait reads it."""
    query = result.table_model.entry_query
    wanted = (ref[0].lower(), ref[1].lower())
    filters = query.filters if query is not None else []
    if not any(isinstance(f, EqualityFilter) and (f.column.table.lower(), f.column.column.lower()) == wanted
               for f in filters):
        return False
    read = {(t.lower(), f.lower()) for t, f in result.coverage.scenario_fields + result.coverage.contact_fields}
    read |= {(a.wait_field[0].lower(), a.wait_field[1].lower()) for a in result.package.journey.activities.values()
             if a.wait_field is not None}
    return wanted not in read


def _how_to_run(result: GenerationResult, context: _TestbookContext) -> str:
    """Spec: "The tester imports the entry CSV directly, together with the CSVs of the linked tables";
    "One scenario at a time"; "Data is read live". The user's decision: the entry table is filled by the
    direct import only, and no automation runs during the test."""
    model = result.table_model
    entry_csv = table_file_name(model.entry_table)
    others = [name for name in context.import_names if name != entry_csv]
    automation = _entry_procedures(result)
    kind = context.entry_kind
    parts = ["### Come eseguire i test"]
    if kind != ENTRY_AUTOMATION:
        parts.extend(_entry_source_paragraphs(result, context, entry_csv, others))
    elif others:
        parts.append(f"Importa prima {_join_words(others)}, poi {entry_csv}, così ogni riga che serve a un contatto è "
                     "già presente quando entra.")
    else:
        parts.append(f"Importa {entry_csv}.")
    if context.managed:
        parts.append(_managed_text(result, context.managed, short=True))
    deadline = activation_deadline_text(result)
    if deadline:
        parts.append(f"{deadline} Dopo quella data i test possono dare un percorso diverso da quello atteso.")
    if automation:
        parts.append(f"Assunzione da verificare al primo run reale: l'ingresso in questo journey è pianificato "
                     f"dall'automation {automation}. Per una entry source di tipo data extension pianificata da "
                     "un'automation, Salesforce potrebbe richiedere il passo di ingresso nel journey di "
                     "quell'automation perché i contatti importati entrino. docs/assumptions.md descrive questa "
                     "assunzione nelle righe da 23 a 27.")
    if context.one_active and context.first_scenario is not None:
        table, field = context.activation_ref
        key_column = _scenario_key_column(context, table)
        row = f" La riga di uno scenario è quella in cui {key_column} è la chiave dello scenario." if key_column else ""
        parts.append("Esegui uno scenario alla volta, in ordine di esecuzione. Gli scenari con un'azione manuale "
                     "vengono per primi, perché la loro scadenza è una data e un'ora fisse. Per rendere attivo uno "
                     f"scenario, imposta {field} in {table} a True sulla riga dello scenario e a False su tutte le "
                     f"altre righe.{row}")
        parts.append(_numbered([
            f"I CSV partono con {context.first_scenario.key} come scenario attivo, quindi il primo run non richiede "
            "modifiche.",
            "Aspetta che i contatti di uno scenario siano usciti dal journey.",
            "Poi rendi attivo lo scenario successivo. Non cambiare altro e non reimportare nulla, perché il journey "
            "legge queste righe mentre i contatti sono al suo interno.",
        ]))
        if automation and _activation_is_entry_filter(result, context.activation_ref):
            parts.append(f"{field} è un filtro della entry query e nessuno split lo legge, quindi con l'import "
                         "diretto non trattiene alcuno scenario.")
    elif kind == ENTRY_RUN_ONCE and result.coverage.scenarios:
        parts.append("Tutti gli scenari entrano insieme quando il journey viene attivato, quindi vengono eseguiti "
                     "insieme: con Run Once non si può eseguire uno scenario alla volta.")
    elif result.coverage.scenarios:
        parts.append("Nei CSV tutti gli scenari sono attivi, quindi vengono eseguiti insieme.")
    mode = context.entry_mode
    if mode == "OnceAndDone":
        parts.append("La modalità di ingresso del journey è OnceAndDone, letta come nessun rientro (assunzione da "
                     "verificare): una chiave del contatto entra una sola volta. Ogni nuovo tentativo richiede quindi "
                     "un nuovo run identifier.")
    elif mode:
        parts.append(f"La modalità di ingresso del journey è {mode}: i contatti di un tentativo precedente potrebbero "
                     "non rientrare. Ogni nuovo tentativo richiede quindi un nuovo run identifier.")
    not_modelled = _not_modelled_text(result.package)
    if not_modelled:
        parts.append(not_modelled)
    parts.extend(_contact_address_text(result))
    parts.extend(_sms_text(result))
    for case in result.coverage.cases:
        smallest = _smallest_branch(case)
        if case.random is None or smallest is None:
            continue
        journey = result.package.journey
        parts.append(f"Il caso statistico {case.id} ha {len(case.contact_keys)} contatti per il random split "
                     f"{_activity_label(journey, case.random.activity_key)}. Ogni contatto prende un branch a caso. "
                     f"L'intervallo accettato è da {case.random.expected_low} a {case.random.expected_high} contatti "
                     f"sul branch {_branch_label(journey, case.random.activity_key, smallest)}.")
    return "\n\n".join(parts)


def _entry_source_paragraphs(result: GenerationResult, context: _TestbookContext, entry_csv: str,
                             others: list[str]) -> list[str]:
    """Spec, "Entry source and procedure", for every entry source other than the automation of the pilot."""
    entry = result.table_model.entry_table
    event = result.package.entry_event
    first = f"Importa prima {_join_words(others)}, poi {entry_csv}" if others else f"Importa {entry_csv}"
    if context.entry_kind == ENTRY_RUN_ONCE:
        return [
            f"L'entry source di questo journey è la data extension {entry} con pianificazione Run Once. {first}, e "
            "solo dopo attiva il journey. Entrano solo le righe presenti nella data extension quando il journey viene "
            "attivato: le righe importate dopo non entrano mai.",
            f"Quando attivi il journey, {entry} deve contenere solo righe di test: ogni riga presente in quel momento "
            "entra nel journey.",
            "Un nuovo tentativo richiede una nuova versione del journey e un nuovo run identifier: genera di nuovo i "
            "dati con un run identifier diverso, importali e attiva la nuova versione.",
        ]
    if context.entry_kind == ENTRY_SCHEDULED:
        mode = getattr(event, "schedule_mode", None)
        schedule = f"una pianificazione ricorrente ({mode})" if mode else "una pianificazione ricorrente non indicata"
        return [
            f"L'entry source di questo journey è la data extension {entry} con {schedule}. {first}. Procedura da "
            "verificare: i contatti dovrebbero entrare alla valutazione successiva della data extension, non al "
            "momento dell'import.",
        ]
    where = "nella sua data extension" if len(others) == 1 else "nelle loro data extension"
    import_text = (f"Importa prima {_join_words(others)} {where}, poi {entry_csv} nella data extension di ingresso "
                   "della copia" if others else f"Importa {entry_csv} nella data extension di ingresso della copia")
    return [
        f"L'entry source di questo journey è {entry_source_text(result.package)}: i contatti non entrano con l'import "
        f"di un CSV. Il test si fa su una copia del journey la cui entry source è una data extension con le stesse "
        f"colonne di {entry}. {import_text}. Poi fai entrare i contatti secondo la pianificazione della copia: con "
        "Run Once, per esempio, importa tutti i CSV prima di attivarla.",
    ]


def _sms_text(result: GenerationResult) -> list[str]:
    """Spec, step 3, "SMS sends": the real number goes to one contact of each case that sends an SMS; without it
    "the SMS cannot be checked and the report says so"."""
    sms_cases = [case for case in result.coverage.cases if getattr(case, "sms_sent", None)]
    if not sms_cases:
        return []
    mobile = getattr(result.package.journey, "default_mobile_field", None)
    where = f" nella colonna {mobile} di {result.table_model.entry_table}" if mobile else ""
    phone = (result.settings.real_phone or "").strip()
    ids = [case.id for case in sms_cases]
    who = f"Il caso {ids[0]} invia" if len(ids) == 1 else f"I casi {_join_words(ids)} inviano"
    if phone:
        return [f"{who} un SMS. Un contatto per caso riceve il numero reale {phone}{where}, e controlla l'SMS sul suo "
                f"telefono; gli altri contatti ricevono il numero fittizio {result.settings.fake_phone}."]
    return [f"{who} un SMS, ma non è stato indicato un numero di telefono reale: l'invio dell'SMS non può essere "
            f"verificato. I contatti ricevono il numero fittizio {result.settings.fake_phone}{where}; l'invio avviene "
            "comunque sul percorso."]


def _contact_address_text(result: GenerationResult) -> list[str]:
    """Spec, "Email": the real inbox with plus-addressing for contacts that must open or click; an
    undeliverable address for all others."""
    kinds = _unique(action.kind for case in result.coverage.cases for action in case.manual_actions)
    actions = _join_words(_ACTION_INFINITIVE.get(kind, kind) for kind in kinds).replace(" e ", " o ") if kinds else ""
    inbox = result.settings.real_inbox.strip()
    real = [contact for contact in result.data.contacts if contact.uses_real_inbox]
    undeliverable = f"un indirizzo non recapitabile che termina con @{UNDELIVERABLE_DOMAIN}"
    if kinds and inbox and real:
        return [f"I contatti che devono {actions} un'email ricevono l'inbox reale {inbox} con plus-addressing, per "
                f"esempio {real[0].email}. Tutti gli altri contatti ricevono {undeliverable}."]
    if kinds:
        return [f"Non è stata indicata un'inbox reale. I contatti che devono {actions} un'email ricevono "
                f"{undeliverable}, quindi quei branch non sono coperti."]
    if result.data.contacts:
        return [f"Nessun contatto deve aprire o cliccare un'email, quindi tutti i contatti ricevono {undeliverable}."]
    return []


def _section_scenarios(result: GenerationResult) -> str:
    journey = result.package.journey
    package = result.package
    context = _testbook_context(result)
    one_table = None
    scenario_rows = []
    for scenario in _ordered_scenarios(result):
        table, items = _values_parts(package, scenario.values, result.coverage.scenario_fields)
        one_table = one_table or table
        keys = [key for case in scenario.cases for key in case.contact_keys]
        active = not context.one_active or scenario is context.first_scenario
        scenario_rows.append([scenario.key, scenario.run_order, "sì" if active else "no",
                              str(bool(scenario.has_manual_action)), str(bool(scenario.is_statistical)),
                              len(scenario.cases), len(keys),
                              f"da {keys[0] if keys else '-'} a {keys[-1] if keys else '-'}",
                              "; ".join(items) or "nessuno"])
    list_intro = "Gli scenari in ordine di esecuzione, come in scenarios.csv."
    if one_table:
        list_intro += f" I valori sono campi di {one_table}."
    case_rows = []
    for scenario in _ordered_scenarios(result):
        for case in scenario.cases:
            if case.random is not None:
                smallest = _smallest_branch(case)
                label = _branch_label(journey, case.random.activity_key, smallest) if smallest else "?"
                special = (f"random split {_activity_label(journey, case.random.activity_key)}: "
                           f"da {case.random.expected_low} a {case.random.expected_high} contatti attesi sul "
                           f"branch {label}")
            else:
                actions = []
                for action in case.manual_actions:
                    deadline = _deadline(context, scenario, case, action)
                    relative = _relative_window(case, action)
                    if deadline:
                        when = f"prima del {deadline}"
                    elif relative is not None:
                        when = f"entro {_minutes_words(relative.minutes)} dall'arrivo"
                    else:
                        when = "appena arriva"
                    actions.append(f"{action.kind} di '{action.email_name}' {when}")
                special = "; ".join(actions)
            if case.random is None and case.unknown_from is not None:
                last = f"non nota da {_activity_label(journey, case.unknown_from)} in poi"
            else:
                last = _activity_label(journey, case.end_activity_key)
            case_rows.append([case.id, scenario.key, len(case.contact_keys), _key_range(case.contact_keys),
                              special or "nessuna", last, _case_summary(context, case)])
    cases_intro = "Un caso è un gruppo di contatti con lo stesso risultato atteso."
    if TESTBOOK_FILE in result.files:
        cases_intro += f" Ogni caso è un test di {TESTBOOK_FILE}, con lo stesso nome."
    return "\n\n".join([
        "## Scenari",
        _how_to_run(result, context),
        "### Elenco degli scenari",
        list_intro,
        _md_table(["Scenario", "Ordine di esecuzione", "Attivo all'inizio", "Azione manuale", "Statistico", "Casi",
                   "Contatti", "Chiavi dei contatti", "Valori"], scenario_rows),
        "### Casi",
        cases_intro,
        _md_table(["Caso", "Scenario", "Contatti", "Chiavi dei contatti", "Azione manuale o random split",
                   "Ultima attività", "Sintesi"], case_rows),
    ])


def _root_text(journey: Journey, root: tuple[str, str], branches: list[BranchStatus],
               by_ref: dict) -> str:
    """One line per root cause: "N branch si trovano a valle di 'split: outcome' (KEY), che non è raggiungibile
    perché ..."."""
    activity = journey.activities.get(root[0])
    name = (activity.name or root[0]) if activity is not None else root[0]
    label = _outcome_label(journey, root[0], root[1])
    status = by_ref.get(root)
    reason = status.reason if status is not None and status.reason else ""
    if not reason:
        inherited = _REASON_BEHIND_BRANCH.match(branches[0].reason or "")
        reason = inherited.group("reason") if inherited else ""
    count = _lie_behind(len(branches))
    if activity is not None and not is_split(activity):
        # Not a split (an unsupported activity, an Update contact whose insert may fail): the simulated paths
        # stop there, so there is no branch to name.
        unsupported = _REASON_UNSUPPORTED.match(reason)
        detail = unsupported.group("detail") if unsupported else reason
        return f"{count} a valle di '{name}' ({root[0]}), dove il percorso simulato si ferma: {detail.rstrip('.')}."
    state, clause = _explain_reason(reason, journey, root[0], root[1])
    if state is None:
        tail = f"che non è coperto: {clause}" if clause else "che non è coperto"
    else:
        tail = f"che {_STATE_WORDS.get(state, state)} perché {clause}"
    return f"{count} a valle di '{name}: {label}' ({root[0]}), {tail.rstrip('.')}."


def _behind_list(journey: Journey, branches: list[BranchStatus]) -> str:
    """The branches behind one root cause, in a compact form: the splits they belong to, by name."""
    outcomes: dict[str, list[str]] = {}
    for branch in branches:
        outcomes.setdefault(branch.activity_key, []).append(branch.outcome_key)
    by_name: dict[str, list[str]] = {}
    complete = True
    for key, keys in outcomes.items():
        activity = journey.activities.get(key)
        entry = key
        if activity is not None and set(keys) != {o.key for o in activity.outcomes}:
            complete = False
            entry = f"{key} per {_join_words(repr(_outcome_label(journey, key, o)) for o in keys)}"
        name = (activity.name or key) if activity is not None else key
        by_name.setdefault(name, []).append(entry)
    count = len(outcomes)
    which = "questo split" if count == 1 else f"questi {count} split"
    lead = f"Sono tutti i branch di {which}:" if complete else f"Sono branch di {which}:"
    items = []
    for name, keys in by_name.items():
        keys = sorted(keys, key=_natural_key)
        items.append(", ".join(keys) if name in keys else f"{name}: {', '.join(keys)}")
    return lead + "\n\n" + _bullets(items)


def _section_not_covered(result: GenerationResult) -> str:
    journey = result.package.journey
    missing = [b for b in result.coverage.branches if b.status == STATUS_NOT_COVERED]
    if not missing:
        return "## Branch non coperti\n\nTutti i branch sono coperti."
    by_ref = {(b.activity_key, b.outcome_key): b for b in result.coverage.branches}
    behind_branch: dict[tuple[str, str], list[BranchStatus]] = {}
    behind_unsupported: dict[str, list[BranchStatus]] = {}
    own: list[BranchStatus] = []
    for branch in missing:
        if branch.behind is not None:
            behind_branch.setdefault(tuple(branch.behind), []).append(branch)
        elif _REASON_BEHIND_UNSUPPORTED.match(branch.reason or ""):
            behind_unsupported.setdefault(branch.reason, []).append(branch)
        else:
            own.append(branch)
    grouped = sum(len(v) for v in behind_branch.values()) + sum(len(v) for v in behind_unsupported.values())
    causes = []
    if behind_branch:
        causes.append(_plural(len(behind_branch), "branch che nessun contatto può raggiungere",
                              "branch che nessun contatto può raggiungere"))
    if behind_unsupported:
        causes.append(_plural(len(behind_unsupported), "elemento non supportato", "elementi non supportati"))
    intro = f"{_plural(len(missing), 'branch non è coperto', 'branch non sono coperti')}."
    if grouped and own:
        intro += (f" Di questi, {grouped} {'si trova' if grouped == 1 else 'si trovano'} a valle di "
                  f"{' o '.join(causes)}: le cause sono spiegate per prime. "
                  f"{'Il restante ha un motivo proprio' if len(own) == 1 else f'I restanti {len(own)} hanno ciascuno un motivo proprio'}, "
                  "indicato nella tabella alla fine.")
    elif grouped:
        intro += (f" {'Si trova' if grouped == 1 else 'Si trovano tutti'} a valle di {' o '.join(causes)}, come "
                  "spiegato di seguito.")
    parts = ["## Branch non coperti", intro]
    if grouped:
        parts.append("### Cause a monte")
        for root, branches in sorted(behind_branch.items(),
                                     key=lambda item: (-len(item[1]), _natural_key(item[0][0]), item[0][1])):
            parts.append(_root_text(journey, root, branches, by_ref))
            parts.append(_behind_list(journey, branches))
        for reason, branches in behind_unsupported.items():
            detail = _REASON_BEHIND_UNSUPPORTED.match(reason).group("detail")
            after = _REASON_AFTER_UNSUPPORTED.match(detail)
            count = len(branches)
            if after:
                text = (f"{count} branch {'viene' if count == 1 else 'vengono'} dopo il branch non supportato "
                        f"'{after.group('label')}' dello stesso split, quindi non "
                        f"{'viene mai valutato' if count == 1 else 'vengono mai valutati'}.")
            else:
                text = f"{_lie_behind(count)} a valle di un elemento non supportato: {detail.rstrip('.')}."
            parts.append(text + " Vedi la sezione Elementi non supportati.")
            parts.append(_behind_list(journey, branches))
    if own:
        rows = []
        for b in own:
            activity = f"{b.activity_name} ({b.activity_key})" if b.activity_name else b.activity_key
            outcome = b.outcome_label if b.outcome_label == b.outcome_key else f"{b.outcome_label} ({b.outcome_key})"
            text = _reason_sentence(b.reason, journey, b.activity_key, b.outcome_key)
            behind = behind_branch.get((b.activity_key, b.outcome_key))
            if behind:
                below = "si trova 1 branch" if len(behind) == 1 else f"si trovano {len(behind)} branch"
                text += f" A valle di questo branch {below}."
            rows.append([activity, outcome, text])
        if grouped:
            parts.append("### Branch con un motivo proprio")
        parts.append(_md_table(["Attività", "Branch", "Motivo"], rows))
    return "\n\n".join(parts)


def _exit_target_text(target: Optional[BranchStatus]) -> str:
    """The state of the coverage target of an exit criterion: the cases that leave that way, or why none does."""
    if target is None:
        return "nessun target"
    if target.status == STATUS_NOT_COVERED:
        return f"non coperto: {_one_line(target.reason).rstrip('.') or 'nessun contatto esce così'}"
    cases = ", ".join(target.covered_by)
    words = "coperto con un'azione manuale" if target.status == STATUS_MANUAL else "coperto"
    return f"{words} dai casi {cases}" if len(target.covered_by) > 1 else f"{words} dal caso {cases}"


def _section_exits(result: GenerationResult) -> str:
    """Spec, step 3, "Exit criteria and goals": every exit criterion with its description, whether the generator
    evaluates it (at the end of every wait) or not and why, and the state of its coverage target; the branches no
    contact can reach because of them are a finding on the journey."""
    journey = result.package.journey
    exits = list(getattr(journey, "exits", None) or [])
    if not exits:
        return ""
    targets = {t.activity_key: t for t in getattr(result.coverage, "exit_targets", None) or []}
    modelled = [e for e in exits if e.unsupported_reason is None]
    text = (f"Il journey ha {_plural(len(exits), 'criterio di uscita', 'criteri di uscita')} (exit criteria). "
            "Journey Builder valuta i criteri di uscita di ogni contatto quando il suo tempo in un wait scade "
            "(documentato).")
    if modelled:
        text += (" Il generatore valuta alla fine di ogni wait del percorso, con l'orologio, i criteri che sa valutare: "
                 "il contatto che li soddisfa esce dal journey in quel punto e il percorso atteso del suo caso finisce "
                 "lì. Ogni criterio è un target di coverage a parte, \"criteri di uscita soddisfatti\".")
    if len(modelled) < len(exits):
        text += (" I criteri non modellati non sono valutati, quindi alcuni contatti potrebbero uscire dal journey "
                 "alla fine di un wait, prima della fine del percorso atteso.")
    rows = [[f"{e.key}: {exit_description(e)}", exit_evaluation_text(e),
             _exit_target_text(targets.get(e.key)) if e.unsupported_reason is None else "non applicabile"]
            for e in exits]
    parts = ["## Criteri di uscita", text,
             _md_table(["Criterio", "Valutazione", "Target criteri di uscita soddisfatti"], rows)]
    behind = [b for b in result.coverage.branches
              if b.status == STATUS_NOT_COVERED and _EXIT_WORDS in (b.reason or "")]
    findings = [w for w in result.coverage.warnings if w.startswith("rilievo sul journey")]
    if behind or findings:
        parts.append(f"{_plural(len(behind), 'branch', 'branch')} non "
                     f"{'è raggiungibile' if len(behind) == 1 else 'sono raggiungibili'} perché i contatti che li "
                     "raggiungerebbero escono prima dal journey per i criteri di uscita. La sezione Branch non coperti "
                     "ne spiega il motivo." if behind else "")
        parts.extend(_sentence(finding) for finding in findings)
    return "\n\n".join(part for part in parts if part)


def _section_special_branches(result: GenerationResult) -> str:
    """Spec, "Outputs": every branch ends in one of three states, each with its reason. The
    branches covered with a manual action, and those crossed only through a random split, are
    listed here by name, with the cases that cover them."""
    journey = result.package.journey
    manual = [b for b in result.coverage.branches if b.status == STATUS_MANUAL]
    statistical = [b for b in result.coverage.branches if b.statistical and b.status != STATUS_NOT_COVERED]
    if not manual and not statistical:
        return ""
    parts = ["## Branch coperti con un'azione manuale o tramite un random split"]
    if manual:
        rows = [[f"{_activity_name(journey, b.activity_key)} ({b.activity_key})", b.outcome_label,
                 ", ".join(b.covered_by)] for b in manual]
        covered = "è coperto" if len(manual) == 1 else "sono coperti"
        parts.append(f"{len(manual)} branch {covered} solo da contatti che devono aprire o cliccare un'email. Il "
                     "testbook indica l'azione e la sua scadenza.")
        parts.append(_md_table(["Attività", "Branch", "Casi"], rows))
    if statistical:
        rows = [[f"{_activity_name(journey, b.activity_key)} ({b.activity_key})", b.outcome_label,
                 ", ".join(b.covered_by)] for b in statistical]
        crossed = "viene attraversato" if len(statistical) == 1 else "vengono attraversati"
        parts.append(f"{len(statistical)} branch {crossed} solo dopo un random split, dai contatti di un caso "
                     "statistico. Il testbook indica l'intervallo accettato.")
        parts.append(_md_table(["Attività", "Branch", "Casi"], rows))
    return "\n\n".join(parts)


def _section_unsupported(result: GenerationResult) -> str:
    journey = result.package.journey
    items = []
    for key in _unique(list(result.package.unsupported_activities) + list(result.coverage.unsupported)):
        activity = journey.activities.get(key)
        label = _activity_label(journey, key)
        reason = getattr(activity, "unsupported_reason", None) if activity is not None else None
        stopped = (getattr(result.coverage, "unsupported_reasons", None) or {}).get(key)
        split_key, _, outcome_key = key.partition("/")
        split = journey.activities.get(split_key) if activity is None and outcome_key else None
        outcome = split.outcome(outcome_key) if split is not None else None
        if stopped and outcome is not None:
            # A branch whose condition orders a field that is not a number or a date: no value inside its range.
            items.append(f"Il branch '{outcome.label or outcome.key}' dello split {_activity_label(journey, split_key)} "
                         f"non è supportato: {stopped.rstrip('.')}. Questo branch e quelli a valle non sono coperti; gli "
                         "altri branch dello split sono valutati.")
        elif stopped:
            # A supported activity the walk cannot predict for some contacts (an insert that may fail).
            why = stopped.rstrip(".")
            items.append(f"L'attività {label} non è supportata per una parte dei contatti: {why}. "
                         "Per questi contatti il percorso non è noto da quell'attività in poi, e i branch che solo "
                         "loro raggiungerebbero non sono coperti.")
        elif outcome is not None:
            # An unsupported branch, written <split>/<outcome> in CoverageResult.unsupported.
            why = (outcome.unsupported_reason or "").rstrip(".")
            items.append(f"Il branch '{outcome.label or outcome.key}' dello split {_activity_label(journey, split_key)} "
                         f"non è supportato{': ' + why if why else ''}. Questo branch, i branch che lo seguono nello "
                         "split e quelli a valle non sono coperti.")
        elif reason:
            items.append(f"L'attività {label} non è supportata: {reason.rstrip('.')}. I suoi branch e quelli a valle "
                         "non sono coperti.")
        elif activity is not None:
            items.append(f"L'attività {label} è di tipo {activity.raw_type}, che il generatore non supporta. I suoi "
                         "branch e quelli a valle non sono coperti.")
        else:
            items.append(f"L'attività {key} non è supportata. I branch a valle non sono coperti.")
    query = result.table_model.entry_query
    if query is not None:
        for flt in query.filters:
            if isinstance(flt, UnsupportedFilter):
                items.append(f"Il filtro `{flt.text}` della entry query non è supportato, quindi le righe generate "
                             "potrebbero non soddisfarlo.")
    items.extend(f"Placeholder: {_sentence(text)}" for text in result.data.placeholders)
    if not items:
        return "## Elementi non supportati\n\nNon è stato trovato alcun elemento non supportato."
    return "\n\n".join(["## Elementi non supportati", _bullets(items)])


def _key_list(text: str) -> list[str]:
    """Activity keys written as ``A, B, C``, in natural order."""
    return sorted((key.strip() for key in text.split(",") if key.strip()), key=_natural_key)


def _case_line(matches: list[re.Match], tail: str, package: Package) -> str:
    """Fields compared with one value written in several letter cases."""
    values = _join_words(_quoted_values(matches[0].group("values"))) or matches[0].group("values")
    refs = _ordered_refs(package, [(m.group("table"), m.group("field")) for m in matches])
    if len(refs) == 1:
        text = f"{refs[0][0]}.{refs[0][1]} viene confrontato con {values}."
    elif len({table.lower() for table, _ in refs}) == 1:
        text = f"{_join_words(field for _, field in refs)} di {refs[0][0]} vengono confrontati ciascuno con {values}."
    else:
        text = f"{_join_words(f'{table}.{field}' for table, field in refs)} vengono confrontati ciascuno con {values}."
    if tail:
        text += f" {_CASE_TAILS.get(tail, _sentence(tail))} Come li confronta Salesforce è un'assunzione da verificare."
    return text


def _same_name_line(matches: list[re.Match]) -> str:
    """Splits that share a name but test different conditions."""
    groups = [f"'{m.group('name')}' in {m.group('count')} versioni, comparse per la prima volta in "
              f"{_join_words(_key_list(m.group('keys')))}" for m in matches]
    return "Alcuni split con lo stesso nome verificano condizioni diverse: " + "; ".join(groups) + "."


def _constant_key_line(matches: list[re.Match]) -> str:
    """Update contact steps that write a constant into a primary-key field."""
    steps = [_key_list(m.group("keys")) for m in matches]
    tables = {m.group("table") for m in matches}
    if all(keys == steps[0] for keys in steps) and len(tables) == 1:
        first = matches[0]
        if len(matches) == 1:
            what = (f"la costante {first.group('value')} nel campo di chiave primaria "
                    f"{first.group('table')}.{first.group('field')}")
        else:
            writes = _join_words(f"{m.group('value')} in {m.group('field')}" for m in matches)
            what = f"costanti nei campi di chiave primaria di {first.group('table')}: {writes}"
        count = _plural(len(steps[0]), "attività Update contact scrive", "attività Update contact scrivono")
        same = "lo stesso valore in quel campo" if len(matches) == 1 else "gli stessi valori in quei campi"
        which = "L'attività è" if len(steps[0]) == 1 else "Le attività sono"
        return f"{count} {what}. Quindi tutti i contatti hanno {same}. {which} {_join_words(steps[0])}."
    parts = [f"{m.group('value')} in {m.group('table')}.{m.group('field')}, da {_join_words(keys)}"
             for m, keys in zip(matches, steps)]
    return ("Alcune attività Update contact scrivono costanti in campi di chiave primaria, quindi tutti i contatti "
            "hanno lo stesso valore in quei campi: " + "; ".join(parts) + ".")


def _one_to_one_line(matches: list[re.Match]) -> str:
    """One-to-one relationships that join on part of a primary key."""
    parts = [f"'{m.group('name')}' collega {m.group('table')} tramite {m.group('fields')}, ma la chiave primaria di "
             f"{m.group('table')} comprende anche {_join_words(k.strip() for k in m.group('extra').split(','))}"
             for m in matches]
    lead = ("Una relazione uno a uno collega una tabella" if len(parts) == 1
            else "Alcune relazioni uno a uno collegano una tabella")
    return (f"{lead} su una parte della sua chiave primaria, quindi a un contatto possono corrispondere più righe: "
            + "; ".join(parts) + ".")


def _merge_patterns(patterns: list[str], package: Package) -> list[str]:
    """One line per kind of suspicious pattern, in the order first met. Patterns of a kind the report does
    not know are kept as they are."""
    kinds = [("case", _CASE_PATTERN), ("names", _SAME_NAME_PATTERN), ("constants", _CONSTANT_KEY_PATTERN),
             ("one_to_one", _ONE_TO_ONE_PATTERN)]
    order: list = []
    groups: dict = {}
    for pattern in patterns:
        for kind, regex in kinds:
            match = regex.match(pattern)
            if match is None:
                continue
            group = ((kind, frozenset(_quoted_values(match.group("values"))), match.group("tail") or "")
                     if kind == "case" else (kind,))
            if group not in groups:
                groups[group] = []
                order.append(group)
            groups[group].append(match)
            break
        else:
            order.append(pattern)
    lines = []
    for item in order:
        if isinstance(item, str):
            lines.append(item)
        elif item[0] == "case":
            lines.append(_case_line(groups[item], item[2], package))
        elif item[0] == "names":
            lines.append(_same_name_line(groups[item]))
        elif item[0] == "constants":
            lines.append(_constant_key_line(groups[item]))
        else:
            lines.append(_one_to_one_line(groups[item]))
    return lines


def _reword_warning(warning: str, result: GenerationResult) -> str:
    """The extra columns of the entry query: what the generator does is a fact, what Salesforce does is an
    assumption (docs/assumptions.md, row 28)."""
    match = _IGNORED_COLUMNS.match(warning.strip())
    if not match:
        return warning
    columns = [column.strip() for column in match.group("columns").split(",") if column.strip()]
    return (f"La entry query seleziona {_plural(len(columns), 'colonna', 'colonne')} che {match.group('table')} non "
            f"ha: {_join_words(columns)}. Il generatore {'la' if len(columns) == 1 else 'le'} ignora. Cosa fa "
            "Salesforce con una query del genere è un'assunzione da verificare, elencata nella sezione Assunzioni.")


def _section_warnings(result: GenerationResult) -> str:
    query = result.table_model.entry_query
    sources = [
        ("Lettura dell'export", result.package.warnings),
        ("Selezione delle tabelle e SQL", list(result.table_model.warnings) + (list(query.warnings) if query else [])),
        ("Coverage", result.coverage.warnings),
        ("Generazione dei dati", result.data.warnings),
        ("Valori che dipendono dall'ora di ingresso", list(getattr(result.coverage, "perishable", None) or [])),
    ]
    patterns: list[str] = []
    groups: list[tuple[str, list[str]]] = []
    for title, warnings in sources:
        found, others = _split_suspicious(warnings)
        patterns.extend(found)
        groups.append((title, [_reword_warning(warning, result) for warning in others]))
    patterns = _merge_patterns(_unique(patterns), result.package)
    parts = ["## Warning e pattern sospetti"]
    for title, items in [("Pattern sospetti", patterns)] + groups:
        items = _unique(_sentence(item) for item in items)
        if items:
            parts.append(f"### {title}\n\n{_bullets(items)}")
    if len(parts) == 1:
        parts.append("Nessun warning.")
    return "\n\n".join(parts)


def _ignored_entry_columns(result: GenerationResult) -> list[str]:
    """Columns that the entry query selects but the entry table does not have."""
    query = result.table_model.entry_query
    entry = _table_of(result.package, result.table_model.entry_table)
    if query is None or entry is None or query.select_all:
        return []
    return [item.target_column for item in query.select if entry.field(item.target_column) is None]


# Text comparisons whose answer depends on letter case (spec assumption "Text comparisons").
_TEXT_TYPES = ("text", "emailaddress", "phone", "locale")
_CASE_OPERATORS = ("equal", "notequal", "is", "isnot", "contains", "beginswith", "endswith")


def _conditions(journey: Journey) -> list[tuple[Activity, Outcome, object]]:
    """Every condition of a decision split outcome the generator reads (unsupported outcomes left out)."""
    return [(activity, outcome, condition) for activity in journey.activities.values() if activity.kind == KIND_DECISION
            for outcome in activity.outcomes if not outcome.unsupported_reason
            for condition in formula_conditions(outcome.formula)]


def _csv_field_types(result: GenerationResult) -> set[str]:
    """The field types of the tables that get a CSV, in lower case."""
    types = set()
    for name in result.table_model.selected_tables:
        table = _table_of(result.package, name)
        types.update((f.type or "").lower() for f in (table.fields if table is not None else []))
    return types


def _linked_by_group_only(result: GenerationResult) -> list[str]:
    """Tables with a CSV linked to the contact only by a group-to-set relationship (no relationship to another
    table): the generator writes the contact key in the relationship's field."""
    selected = {name.lower() for name in result.table_model.selected_tables}
    relationships = list(result.package.relationships)
    between_tables = {name.lower() for r in relationships if r.left_table is not None
                      for name in (r.left_table, r.right_table)}
    return _unique(r.right_table for r in relationships if r.left_table is None
                   and r.right_table.lower() in selected and r.right_table.lower() not in between_tables)


def _one_to_many(result: GenerationResult) -> list[str]:
    """Tables with a CSV on the "many" side of a relationship (right cardinality Many: several rows of the right
    table for one row of the left table, or for one contact when the left side is the contact)."""
    selected = {name.lower() for name in result.table_model.selected_tables}
    return _unique(r.right_table for r in result.package.relationships
                   if (r.right_cardinality or "").lower() == "many" and r.right_table.lower() in selected)


def _triggered_assumptions(result: GenerationResult) -> list[str]:
    """Spec, "Honest outcomes": "report.md lists only the assumptions the export triggers". Each item names an
    assumption of docs/assumptions.md that this journey's expected results rest on."""
    settings = result.settings
    journey = result.package.journey
    activities = list(journey.activities.values())
    kinds = {activity.kind for activity in activities}
    conditions = _conditions(journey)
    items = []
    # The six assumptions of the spec ("Assumptions to verify and known limits").
    if any(sum(1 for o in a.outcomes if o.formula is not None and not o.is_remainder) >= 2
           for a in activities if a.kind == KIND_DECISION):
        items.append("I branch di un decision split vengono valutati in ordine e vince il primo la cui condizione è "
                     "soddisfatta.")
    if any(c.value is not None and c.value_attribute is None and c.operator.lower() in _CASE_OPERATORS
           and _field_type(result.package, c.table, c.field) in _TEXT_TYPES for _, _, c in conditions):
        if settings.case_insensitive_text:
            items.append("I confronti di testo non distinguono maiuscole e minuscole, quindi `Email` e `EMAIL` sono "
                         "lo stesso valore. L'impostazione che ignora maiuscole e minuscole è attiva.")
        else:
            items.append("I confronti di testo distinguono maiuscole e minuscole, quindi `Email` e `EMAIL` sono valori "
                         "diversi. L'impostazione che ignora maiuscole e minuscole è disattivata.")
    if KIND_WAIT_ATTRIBUTE in kinds:
        items.append("Un wait by attribute la cui data è già passata termina subito.")
    types = _csv_field_types(result)
    dates = (f"date scritte come {_date(result.generation_time, settings)}, nel formato "
             f"{_human_date_format(settings.date_format)}")
    if "date" in types and "boolean" in types:
        items.append(f"L'import accetta {dates}, e `True` e `False` per i Boolean.")
    elif "date" in types:
        items.append(f"L'import accetta {dates}.")
    elif "boolean" in types:
        items.append("L'import accetta `True` e `False` per i Boolean.")
    if KIND_ENGAGEMENT in kinds:
        items.append("Un engagement split viene valutato quando il contatto lo raggiunge, non quando l'email viene "
                     "inviata.")
    if KIND_WAIT_ATTRIBUTE in kinds:
        zones = sorted({a.wait_timezone for a in activities if a.kind == KIND_WAIT_ATTRIBUTE and a.wait_timezone})
        zone_text = (f" I wait by attribute di questo journey sono impostati sul fuso orario {_join_words(zones)}."
                     if zones else "")
        items.append("Il fuso orario applicato alla data del wait non è noto, quindi le scadenze del testbook "
                     f"potrebbero spostarsi di qualche ora.{zone_text}")
    # The entry source of the pilot.
    automation = _entry_procedures(result)
    if automation:
        items.append(f"L'import diretto di {table_file_name(result.table_model.entry_table)} fa entrare i suoi "
                     f"contatti. L'ingresso in questo journey è pianificato dall'automation {automation} e la review "
                     "Salesforce indica che potrebbe servire il passo di ingresso nel journey di quell'automation. "
                     "Vedi la sezione Come eseguire i test.")
    ignored = _ignored_entry_columns(result)
    if ignored and result.table_model.entry_query is not None:
        items.append(f"La entry query {result.table_model.entry_query.name} viene eseguita anche se seleziona "
                     f"{_plural(len(ignored), 'colonna', 'colonne')} che {result.table_model.entry_table} non ha. "
                     f"Il generatore ignora {'quella colonna' if len(ignored) == 1 else 'quelle colonne'}.")
    items.extend(_extension_assumptions(result, conditions))
    items.extend(_time_assumptions(result, conditions))
    return items


def _time_assumptions(result: GenerationResult, conditions: list) -> list[str]:
    """The assumptions of the time-dependent extensions (spec, step 3) that this export triggers. The day-level
    comparison of relative dates is a fact confirmed by the team (14 October 2026), so it is not listed; the
    comparison of the full date and time with hours or minutes (docs/assumptions.md row 48) still is."""
    journey = result.package.journey
    items = []
    every = list(conditions) + [(None, None, c) for e in getattr(journey, "exits", None) or []
                                for c in formula_conditions(e.formula)]
    for _, _, condition in every:
        relative = getattr(condition, "relative", None)
        try:
            full_time = relative is not None and not criteria.is_day_level(relative)
        except criteria.UnsupportedOperator:
            continue
        if full_time:
            items.append("Una condizione con una data relativa in ore o minuti confronta data e ora intere, non "
                         "solo il giorno.")
            break
    return items


def _extension_assumptions(result: GenerationResult, conditions: list) -> list[str]:
    """The assumptions added with the extensions beyond the pilot (docs/assumptions.md), when this export
    triggers them."""
    journey = result.package.journey
    kind = entry_kind(result.package)
    entry = result.table_model.entry_table
    items = []
    if kind in (ENTRY_RUN_ONCE, ENTRY_SCHEDULED):
        items.append("L'automation che l'entry event indica (automationId), quando c'è, è lo scheduler che Journey "
                     f"Builder crea per la pianificazione della data extension, non un'automation che popola {entry}: "
                     "per questo non c'è una entry query da eseguire.")
    if kind == ENTRY_SCHEDULED:
        items.append("Con una pianificazione ricorrente, i contatti importati entrano alla valutazione successiva "
                     "della data extension di ingresso.")
    if getattr(journey, "entry_mode", None) == "OnceAndDone":
        items.append("La modalità di ingresso OnceAndDone è \"No re-entry\": una chiave del contatto entra nel "
                     "journey una sola volta, in tutte le sue versioni.")
    key = result.table_model.contact_key_field
    key = (key[0].lower(), key[1].lower()) if key else None
    tested = _unique(f"{c.table}.{c.field}" for _, _, c in conditions
                     if c.operator in ("IsNull", "IsNotNull") and (c.table.lower(), c.field.lower()) == key)
    if tested:
        items.append(f"Un contatto senza chiave del contatto non può entrare nel journey, quindi una condizione "
                     f"IsNull o IsNotNull su {_join_words(tested)} dà sempre lo stesso risultato.")
    groups = _linked_by_group_only(result)
    if groups:
        items.append(f"La relazione group-to-set che collega {_join_words(groups)} al contatto usa la Contact Key "
                     "(leftRelationshipIDs 3), quindi il generatore scrive la chiave del contatto nel suo campo.")
    compared = _unique(f"{c.table}.{c.field} è letto come \"il contatto ha una riga in {c.table}\""
                       for _, _, c in conditions if c.value_attribute is not None)
    if compared:
        items.append(f"Il confronto tra attributi su {'; su '.join(compared)}. Salesforce documenta alcuni known issue "
                     "sui confronti tra attributi (campi che ammettono null, gruppi AND e OR annidati): il percorso "
                     "atteso assume che il confronto funzioni.")
    many = _one_to_many(result)
    if many:
        items.append(f"Una condizione su una tabella in relazione uno a molti ({_join_words(many)}) è vera se una "
                     "qualsiasi riga del contatto la soddisfa. Il generatore scrive una sola riga per contatto.")
    row_tables = {table.lower(): table for table in (getattr(result.coverage, "row_fields", None) or {})}
    inserted = _unique(row_tables[w.table.lower()] for activity in journey.activities.values()
                       if activity.kind == KIND_UPDATE for w in activity.writes if w.table.lower() in row_tables)
    if inserted:
        items.append(f"Un Update contact che inserisce la riga mancante del contatto in {_join_words(inserted)} dà ai "
                     "campi che non scrive il loro valore di default, o null se non ne hanno. Se lascia vuoto un campo "
                     "che non ammette null e non ha un valore di default, l'inserimento fallisce e il contatto esce "
                     "dal journey.")
    return items


def _section_assumptions(result: GenerationResult) -> str:
    """Spec, "Assumptions to verify and known limits": the assumptions the generator encodes; spec, "Honest
    outcomes": "report.md lists only the assumptions the export triggers"."""
    items = _triggered_assumptions(result)
    if not items:
        return "\n\n".join(["## Assunzioni", "I risultati attesi di questo journey non dipendono da nessuna delle "
                              "assunzioni elencate nel file docs/assumptions.md del repository CJ Test Helper."])
    return "\n\n".join([
        "## Assunzioni",
        "I risultati attesi dipendono da queste assunzioni sul comportamento di Salesforce. Verifica ciascuna una "
        "volta con un piccolo run reale. Il file docs/assumptions.md del repository CJ Test Helper le elenca tutte, "
        "con il modo di verificarle e quali controllare per prime.",
        _bullets(items),
    ])


_SETTING_LABELS = {
    "contacts_per_case": "Contatti per caso",
    "contacts_per_random_case": "Contatti per caso di random split",
    "real_inbox": "Inbox reale per i test di open e click",
    "real_phone": "Numero di telefono reale",
    "fake_phone": "Numero di telefono fittizio",
    "window_start_days": "Inizio della finestra di validità",
    "window_end_days": "Fine della finestra di validità",
    "wait_offset_hours": "Offset della data di wait",
    "wait_offset_before_engagement_hours": "Offset della data di wait prima di un engagement split",
    "date_format": "Formato delle date nei CSV",
    "one_active_scenario": "Un solo scenario attivo alla volta",
    "activation_field": "Campo di attivazione",
    "run_id": "Run identifier nelle chiavi",
    "case_insensitive_text": "I confronti di testo ignorano maiuscole e minuscole",
    "seed": "Seed del set cover",
    "generation_time": "Ora di generazione",
    "entry_time": "Ora di ingresso",
    "min_validity_days": "Validità minima dei dati (giorni)",
}


def _days_text(days: int) -> str:
    if days == 0:
        return "l'ora di generazione"
    side = "prima" if days < 0 else "dopo"
    return f"{_plural(abs(days), 'giorno', 'giorni')} {side} dell'ora di generazione"


def _setting_text(name: str, value, result: GenerationResult) -> str:
    if name == "generation_time":
        return _date(result.generation_time, result.settings)
    if name == "entry_time":
        return _date(value, result.settings) if value is not None else "l'ora di generazione"
    if name == "min_validity_days" and isinstance(value, int):
        return _plural(value, "giorno", "giorni")
    if name == "one_active_scenario" and value is None:
        return "sì, quando viene rilevato un campo di attivazione"
    if name == "activation_field" and value is None:
        return "suggerito dai filtri della entry query"
    if name in ("window_start_days", "window_end_days") and isinstance(value, int):
        return _days_text(value)
    if name in ("wait_offset_hours", "wait_offset_before_engagement_hours") and isinstance(value, int):
        return _plural(value, "ora", "ore")
    if name == "date_format" and isinstance(value, str):
        return _human_date_format(value)
    if value is None or value == "":
        return "nessuno"
    if isinstance(value, bool):
        return "sì" if value else "no"
    if isinstance(value, tuple):
        return ".".join(str(part) for part in value)
    return str(value)


def _section_settings(result: GenerationResult) -> str:
    rows = [[_SETTING_LABELS.get(f.name, f.name), _setting_text(f.name, getattr(result.settings, f.name), result)]
            for f in dataclasses.fields(result.settings)]
    on, ref, _ = activation_setting(result.package, result.table_model, result.settings)
    if on:
        applied = f"sì, campo di attivazione {ref[0]}.{ref[1]}"
    elif one_active_not_offered(result.package):
        applied = "no: non si applica a un ingresso Run Once, dove tutti gli scenari entrano insieme"
    else:
        applied = "no"
    rows.append(["Un solo scenario attivo alla volta, come applicato", applied])
    return "\n\n".join(["## Impostazioni usate", _md_table(["Impostazione", "Valore"], rows)])


def write_report(result: GenerationResult) -> str:
    """Spec, "Outputs": "report.md: Coverage, unreachable branches, unsupported elements, warnings,
    settings used". It "opens with a summary: number of scenarios, cases and contacts, and branches
    covered, covered with manual action, and not covered". Every number is computed from ``result``."""
    sections = [
        _section_title(result),
        _section_summary(result),
        _section_checks(result),
        _section_tables(result),
        _section_scenarios(result),
        _section_not_covered(result),
        _section_exits(result),
        _section_special_branches(result),
        _section_unsupported(result),
        _section_warnings(result),
        _section_assumptions(result),
        _section_settings(result),
    ]
    return "\n\n".join(sec for sec in sections if sec) + "\n"
